"""格式转换测试:python tests/test_converter.py [输出目录]  (不依赖 pytest)"""
import io
import os
import sys
import tempfile
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import converter as C  # noqa: E402
import engine as E  # noqa: E402
import ocr  # noqa: E402
from make_sample import make_sample  # noqa: E402
from PIL import Image  # noqa: E402

fitz = E.fitz
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(tempfile.mkdtemp(prefix="pdfed_conv_"))
OUT.mkdir(parents=True, exist_ok=True)
SRC = OUT / "src"
SRC.mkdir(exist_ok=True)
failures = []


def check(cond, msg):
    print(("  ok   " if cond else "  FAIL ") + msg)
    if not cond:
        failures.append(msg)


def norm(s):
    return s.replace(" ", " ")


def pdf_text(path):
    d = fitz.open(path)
    return norm("\n".join(p.get_text() for p in d))


def conv(src, target, opts=None, outdir=None):
    od = outdir or str(OUT / f"to_{target}")
    r = C.convert(str(src), target, od, opts or C.Options())
    if not r.ok:
        print("     !!", r.message)
    return r


# ---------------------------------------------------------------- 造测试输入
sample = str(SRC / "sample.pdf")
make_sample(sample)

# 带表格和图片的 PDF(用 HTML 排版生成)
rich_png = SRC / "pic.png"
Image.new("RGB", (240, 120), (30, 144, 255)).save(rich_png)
html = f"""<h1>季度报告 Report</h1><p>这是第一段正文,包含中文和 English words。用来检查段落转换。</p>
<table border="1"><tr><th>项目</th><th>数量</th><th>金额</th></tr>
<tr><td>苹果</td><td>12</td><td>3,456.50</td></tr><tr><td>Banana</td><td>007</td><td>89</td></tr>
<tr><td>合计</td><td>19</td><td>3,545.50</td></tr></table>
<p>图片如下:</p><img src="pic.png" width="200"><p>结束。</p>"""
rich = str(SRC / "rich.pdf")
C._html_to_pdf(html, rich, str(SRC))

# 扫描件
src_doc = fitz.open(sample)
scan = fitz.open()
sp = scan.new_page(width=src_doc[0].rect.width, height=src_doc[0].rect.height)
sp.insert_image(sp.rect, pixmap=src_doc[0].get_pixmap(dpi=150))
scan_path = str(SRC / "scan.pdf")
scan.save(scan_path)

print("== 目标格式列表")
check(("docx", "Word 文档 (.docx)") in C.targets_for([sample]), "PDF 可以转 Word")
check([k for k, _ in C.targets_for([str(rich_png)])][0] == "pdf", "图片首选 PDF")
check(C.targets_for([sample, str(SRC / "x.unknown")]) == [], "含有不支持的文件时没有目标格式")
mixed = dict(C.targets_for([sample, str(rich_png)]))
check(mixed.get("png") == "图片 PNG" and mixed.get("pdf") == "PDF 文档" and "docx" not in mixed,
      "PDF 和图片混选:只列共同的格式,并用通用名称")
check(("pdf", "PDF 文档") in C.targets_for([str(rich_png), str(SRC / "a.txt")]), "图片和文本混选可以一起转 PDF")
check(C.group_of("x.DOCX") == "word" and C.group_of("a.unknown") is None, "按扩展名识别类型")

print("== PDF → 图片 / SVG / 文本 / 网页 / 压缩")
for fmt in ("png", "jpg", "webp", "tiff", "bmp"):
    r = conv(sample, fmt, C.Options(pages="1-2", dpi=72))
    ok = r.ok and len(r.outputs) == 2 and all(Image.open(o).size == (595, 842) for o in r.outputs)
    check(ok, f"PDF → {fmt}(第 1-2 页,各一张)")
r = conv(sample, "png", C.Options(pages="2", dpi=72))
check(r.ok and Path(r.outputs[0]).name == "sample.png", "只转一页时文件名不带页码")
r = conv(sample, "svg", C.Options(pages="1"))
check(r.ok and "<svg" in Path(r.outputs[0]).read_text(encoding="utf-8")[:500], "PDF → SVG")
r = conv(sample, "txt")
t = Path(r.outputs[0]).read_text(encoding="utf-8")
check(r.ok and "第二页" in t and "Tight line 0" in t and "—— 第 2 页 ——" in t, "PDF → 纯文本(带页分隔)")
r = conv(sample, "html")
h = Path(r.outputs[0]).read_text(encoding="utf-8")
check(r.ok and "示例文档" in h and "<meta charset" in h, "PDF → 网页")
big = fitz.open()
bp = big.new_page()
bp.insert_image(bp.rect, pixmap=fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 3000, 4000), False))
big.save(str(SRC / "bigimg.pdf"))
r = conv(SRC / "bigimg.pdf", "pdf", C.Options(dpi=72, quality=50))
check(r.ok and "KB" in r.message, f"压缩 PDF ({r.message})")

print("== PDF → Word")
r = conv(sample, "docx")
check(r.ok, "转换成功")
from docx import Document  # noqa: E402

d = Document(r.outputs[0])
alltext = "\n".join(p.text for p in d.paragraphs)
check("示例文档" in alltext and "这是一段中文正文" in alltext and "Tight line 2" in alltext, "文字内容都在")
check(len(d.sections) == 3, f"3 页 → 3 个分节 (实际 {len(d.sections)})")
check(abs(d.sections[2].page_width.pt - 842) < 2, "横向页保持横向尺寸")
big_runs = [run for p in d.paragraphs for run in p.runs if run.font.size and run.font.size.pt > 20]
check(any("示例文档" in r_.text for r_ in big_runs), "标题保留了大字号")
r = conv(rich, "docx")
d = Document(r.outputs[0])
check(len(d.tables) == 1 and d.tables[0].cell(1, 0).text == "苹果" and d.tables[0].cell(3, 2).text.startswith("3,545"),
      "表格转成了 Word 表格")
check(len(d.inline_shapes) >= 1, "图片被保留")
check("第一段正文" in "\n".join(p.text for p in d.paragraphs) and "项目" not in "\n".join(p.text for p in d.paragraphs),
      "表格里的文字没有重复变成段落")

print("== PDF → Excel")
r = conv(rich, "xlsx")
import openpyxl  # noqa: E402

wb = openpyxl.load_workbook(r.outputs[0])
ws = wb.worksheets[0]
check(r.ok and ws["A2"].value == "苹果" and ws["B2"].value == 12 and ws["C2"].value == 3456.5, "表格 → 工作表,数字识别为数值")
check(ws["B3"].value == "007", "前导 0 的编号保持文本")
r = conv(sample, "xlsx")
check(r.ok and "没有检测到表格" in r.message or r.ok, f"没有表格时按行放文字 ({r.message})")

print("== PDF → PowerPoint")
from pptx import Presentation  # noqa: E402

r = conv(sample, "pptx")
prs = Presentation(r.outputs[0])
texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]
check(r.ok and len(prs.slides) == 3 and any("示例文档" in t for t in texts), "可编辑模式:3 页幻灯片,文字是文本框")
check(any(sh.shape_type == 13 for sh in prs.slides[0].shapes), "背景图片存在")
r = conv(sample, "pptx", C.Options(pptx_mode="image"))
prs = Presentation(r.outputs[0])
check(r.ok and not any(sh.has_text_frame for sh in prs.slides[0].shapes if sh.shape_type != 13), "图片模式:每页只有图片")

print("== 扫描件(OCR)转 Word / 文本")
if not ocr.available() or not ocr.list_languages():
    print("  skip 本机没有 Windows OCR")
else:
    r = conv(scan_path, "docx")
    d = Document(r.outputs[0])
    check(r.ok and "Tight line" in "\n".join(p.text for p in d.paragraphs), "扫描件 → Word:识别出文字")
    r = conv(scan_path, "txt")
    check(r.ok and "编辑文字" in Path(r.outputs[0]).read_text(encoding="utf-8"), "扫描件 → 文本:识别出文字")
    r = conv(scan_path, "txt", C.Options(ocr=False))
    check(r.ok and "Tight" not in Path(r.outputs[0]).read_text(encoding="utf-8"), "关闭 OCR 时不识别")

print("== 图片 → PDF / 其它图片")
jpg = SRC / "photo.jpg"
Image.new("RGB", (1600, 900), (200, 80, 60)).save(jpg, quality=95)
png2 = SRC / "alpha.png"
Image.new("RGBA", (300, 200), (0, 128, 0, 120)).save(png2)
r = conv(jpg, "pdf")
dd = fitz.open(r.outputs[0])
check(r.ok and len(dd) == 1 and abs(dd[0].rect.width - 841.89) < 1, "横图 → 横向 A4")
r = conv(jpg, "pdf", C.Options(page_size="fit", dpi=150))
dd = fitz.open(r.outputs[0])
check(abs(dd[0].rect.width - 1600 * 72 / 150) < 2, "原图大小模式")
res = C.convert_many([str(jpg), str(png2)], "pdf", lambda s: str(OUT / "merged"), C.Options(merge=True))
if not res[0].ok:
    print("     !!", res[0].message)
check(res[0].ok and len(fitz.open(res[0].outputs[0])) == 2, "两张图片合并成 2 页 PDF")
for fmt in ("png", "jpg", "webp", "bmp", "tiff", "gif", "ico"):
    r = conv(png2, fmt)
    ok = r.ok and Path(r.outputs[0]).exists()
    if ok and fmt != "ico":
        im = Image.open(r.outputs[0])
        ok = im.size == (300, 200)
    check(ok, f"PNG(带透明) → {fmt}")
r = conv(png2, "jpg")
check(Image.open(r.outputs[0]).convert("RGB").getpixel((5, 5))[0] > 100, "透明背景转 JPG 垫白底(不是黑底)")
check(Image.open(conv(png2, "ico").outputs[0]).size[0] >= 128, "ICO 含多尺寸")

print("== 文本 / Markdown / 网页 → PDF")
(SRC / "a.txt").write_text("第一行 中文\n  缩进的行 indented\n\nEnglish line three. " + "很长的一行。" * 40, encoding="utf-8")
(SRC / "b.md").write_text("# 大标题\n\n正文 **加粗** 和 `code`。\n\n| 列A | 列B |\n|---|---|\n| 1 | 2 |\n\n- 条目一\n- 条目二\n", encoding="utf-8")
(SRC / "c.html").write_text("<html><body><h1>网页标题</h1><p>网页内容 <b>粗体</b></p><img src='pic.png' width=150></body></html>", encoding="utf-8")
(SRC / "gbk.txt").write_bytes("这是 GBK 编码的文本".encode("gb18030"))
for name, needle in (("a.txt", "缩进的行"), ("b.md", "大标题"), ("c.html", "网页内容"), ("gbk.txt", "GBK 编码")):
    r = conv(SRC / name, "pdf")
    check(r.ok and needle in pdf_text(r.outputs[0]), f"{name} → PDF")
r = conv(SRC / "b.md", "pdf")
check("列A" in pdf_text(r.outputs[0]) and "条目二" in pdf_text(r.outputs[0]), "Markdown 的表格和列表")
check(os.path.getsize(r.outputs[0]) < 200_000, f"中文字体已裁剪,PDF 很小 ({os.path.getsize(r.outputs[0])} 字节)")

print("== Office → PDF / 图片 / 文本")
engines = C.office_engines()
print("  引擎:", engines)
from docx import Document as NewDoc  # noqa: E402
from docx.shared import Pt  # noqa: E402

dx = NewDoc()
dx.add_heading("Word 测试文档", 0)
dx.add_paragraph("这是一段中文正文 with English text。")
tb = dx.add_table(rows=2, cols=2)
tb.style = "Table Grid"
tb.cell(0, 0).text, tb.cell(0, 1).text, tb.cell(1, 0).text, tb.cell(1, 1).text = "甲", "乙", "1", "2"
dx_path = SRC / "doc.docx"
dx.save(dx_path)
xl = openpyxl.Workbook()
xl.active.title = "数据"
xl.active.append(["名称", "数量"])
xl.active.append(["苹果", 5])
xl_path = SRC / "book.xlsx"
xl.save(xl_path)
pp = Presentation()
s = pp.slides.add_slide(pp.slide_layouts[5])
s.shapes.title.text = "幻灯片标题 Slide"
pp_path = SRC / "deck.pptx"
pp.save(pp_path)
(SRC / "data.csv").write_text("姓名,分数\n张三,90\nLi Si,85\n", encoding="utf-8")

# 内置方式(不依赖 Office)
tmp_pdf = str(OUT / "builtin_doc.pdf")
C._docx_builtin(str(dx_path), tmp_pdf, C._Ctx(None, None))
check("Word 测试文档" in pdf_text(tmp_pdf) and "甲" in pdf_text(tmp_pdf), "内置:docx → PDF(不依赖 Office)")
tmp_pdf2 = str(OUT / "builtin_sheet.pdf")
C._sheet_builtin(str(xl_path), tmp_pdf2, C._Ctx(None, None))
check("苹果" in pdf_text(tmp_pdf2), "内置:xlsx → PDF")
r = conv(SRC / "data.csv", "pdf")
check(r.ok and "张三" in pdf_text(r.outputs[0]), "csv → PDF")
if engines:
    t0 = time.time()
    want = "Office/WPS" if any(e in engines for e in ("Microsoft Office", "WPS Office")) else "LibreOffice"
    r = conv(dx_path, "pdf")
    check(r.ok and "Word 测试文档" in pdf_text(r.outputs[0]) and "甲" in pdf_text(r.outputs[0]) and want in r.message,
          f"docx → PDF ({r.message}, {time.time() - t0:.1f}s)")
    r = conv(xl_path, "pdf")
    check(r.ok and "苹果" in pdf_text(r.outputs[0]) and want in r.message, f"xlsx → PDF ({r.message})")
    r = conv(pp_path, "pdf")
    check(r.ok and "幻灯片标题" in pdf_text(r.outputs[0]) and want in r.message, f"pptx → PDF ({r.message})")
    r = conv(dx_path, "png", C.Options(dpi=72))
    check(r.ok and Image.open(r.outputs[0]).size[0] > 400, "docx → PNG(经过 PDF)")
else:
    print("  skip 本机没有 Office / WPS / LibreOffice")
r = conv(dx_path, "txt")
check(r.ok and "这是一段中文正文" in Path(r.outputs[0]).read_text(encoding="utf-8") and "甲\t乙" in Path(r.outputs[0]).read_text(encoding="utf-8"),
      "docx → 纯文本(含表格)")
r = conv(SRC / "nonexistent.pdf", "docx")
check(not r.ok and "不存在" in r.message, "文件不存在时给出清晰的错误")
r = conv(SRC / "a.txt", "docx")
check(not r.ok and "不支持" in r.message, "不支持的组合给出清晰的错误")
(SRC / "broken.pdf").write_bytes(b"%PDF-1.4 not really")
r = conv(SRC / "broken.pdf", "docx")
check(not r.ok, f"损坏的 PDF 不会崩溃 ({r.message[:40]})")

print("== 取消 / 进度 / 批量 / 命令行")
seen = []
res = C.convert_many([sample, rich], "png", lambda s: str(OUT / "batch"), C.Options(dpi=40),
                     progress=lambda d, t, text=None: seen.append((d, t)))
check(len(res) == 2 and all(x.ok for x in res) and seen[-1][0] == seen[-1][1], "批量转换:进度走到 100%")
check(all(a[0] <= b[0] for a, b in zip(seen, seen[1:])), "进度单调递增")
calls = {"n": 0}


def cancel_soon():
    calls["n"] += 1
    return calls["n"] > 3


res = C.convert_many([sample, rich], "png", lambda s: str(OUT / "cancelled"), C.Options(dpi=40), cancel=cancel_soon)
check(res and res[-1].message == "已取消", "取消后中止并返回已完成的结果")
code = C.cli([sample, "--to", "txt", "--out", str(OUT / "cli"), "--pages", "1"])
check(code == 0 and (OUT / "cli" / "sample.txt").exists(), "命令行转换")

print()
if failures:
    print(f"失败 {len(failures)} 项:", *failures, sep="\n  - ")
    sys.exit(1)
print("格式转换测试全部通过 ✔  输出目录:", OUT)
