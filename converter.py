"""文件格式转换。与界面无关;可以在后台线程里调用(同一时间只能有一个线程在用 PyMuPDF)。

支持的转换
----------
PDF   → Word(.docx)、Excel(.xlsx)、PowerPoint(.pptx)、图片(png/jpg/webp/tiff/bmp)、纯文本、网页、SVG、压缩后的 PDF
Word / Excel / PowerPoint → PDF、图片(通过 PDF)。用本机的 Microsoft Office 或 WPS(COM),或 LibreOffice;
                         没有这些软件时,.docx / .xlsx / .csv 还可以用内置方式转成 PDF
图片  → PDF(单张或合并多张)、其它图片格式
txt / md / html → PDF

扫描件(整页是图片的 PDF)在转成 Word / Excel / 文本时会自动先做文字识别(OCR)。
"""
from __future__ import annotations

import html as htmllib
import io
import os
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path

import engine as E
import ocr
from PIL import Image, ImageOps

fitz = E.fitz

# ----------------------------------------------------------------------------
# 类型与目标格式
# ----------------------------------------------------------------------------

GROUPS = {
    "pdf": {".pdf"},
    "word": {".doc", ".docx", ".rtf", ".odt", ".wps"},
    "excel": {".xls", ".xlsx", ".ods", ".csv"},
    "ppt": {".ppt", ".pptx", ".odp"},
    "image": {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".webp", ".ico"},
    "text": {".txt", ".md", ".markdown", ".html", ".htm"},
}
GROUP_NAMES = {"pdf": "PDF", "word": "Word 文档", "excel": "Excel 表格", "ppt": "PowerPoint", "image": "图片",
               "text": "文本 / 网页"}

_IMG_TARGETS = [("png", "图片 PNG"), ("jpg", "图片 JPG"), ("webp", "图片 WebP"), ("tiff", "图片 TIFF"), ("bmp", "图片 BMP")]
TARGETS = {
    "pdf": [("docx", "Word 文档 (.docx)"), ("xlsx", "Excel 表格 (.xlsx)"), ("pptx", "PowerPoint 演示文稿 (.pptx)")]
           + [(k, label + "(每页一张)") for k, label in _IMG_TARGETS]
           + [("txt", "纯文本 (.txt)"), ("html", "网页 (.html)"), ("svg", "SVG 矢量图(每页一个)"),
              ("pdf", "压缩优化后的 PDF")],
    "word": [("pdf", "PDF 文档"), ("png", "图片 PNG(每页一张)"), ("jpg", "图片 JPG(每页一张)"), ("txt", "纯文本 (.txt,仅 docx)")],
    "excel": [("pdf", "PDF 文档"), ("png", "图片 PNG(每页一张)"), ("jpg", "图片 JPG(每页一张)")],
    "ppt": [("pdf", "PDF 文档"), ("png", "图片 PNG(每张幻灯片一张)"), ("jpg", "图片 JPG(每张幻灯片一张)")],
    "image": [("pdf", "PDF 文档")] + _IMG_TARGETS + [("gif", "图片 GIF"), ("ico", "ICO 图标")],
    "text": [("pdf", "PDF 文档")],
}
IMAGE_FORMATS = {"png", "jpg", "webp", "tiff", "bmp", "gif", "ico"}


def group_of(path: str):
    ext = Path(path).suffix.lower()
    for g, exts in GROUPS.items():
        if ext in exts:
            return g
    return None


def all_extensions() -> list:
    return sorted(e for exts in GROUPS.values() for e in exts)


_GENERIC_LABELS = {"pdf": "PDF 文档", "png": "图片 PNG", "jpg": "图片 JPG", "webp": "图片 WebP", "tiff": "图片 TIFF",
                   "bmp": "图片 BMP", "txt": "纯文本 (.txt)"}


def targets_for(paths) -> list:
    """所选文件都能转成的目标格式 [(key, 显示名)]。混选不同类型时用通用名称(各自按自己的方式转换)。"""
    keys, first, groups = None, [], set()
    for p in paths:
        g = group_of(p)
        if g is None:
            return []
        groups.add(g)
        t = TARGETS[g]
        if keys is None:
            keys, first = {k for k, _ in t}, t
        else:
            keys &= {k for k, _ in t}
    out = [(k, lab) for k, lab in first if keys and k in keys]
    if len(groups) > 1:
        out = [(k, _GENERIC_LABELS.get(k, lab)) for k, lab in out]
    return out


@dataclass
class Options:
    pages: str = ""            # 页码范围,空 = 全部,如 "1-3,5"
    dpi: int = 150             # 转图片 / 幻灯片图片的清晰度
    quality: int = 90          # JPG / WebP / 压缩 PDF 的图片质量
    merge: bool = False        # 多张图片合并成一个 PDF
    page_size: str = "a4"      # 图片转 PDF:"a4" 适应 A4,"fit" 原图大小
    ocr: bool = True           # 扫描页自动文字识别
    ocr_lang: str | None = None
    pptx_mode: str = "editable"  # "editable" 可编辑文字 + 背景图 / "image" 每页一张图


@dataclass
class Result:
    src: str
    outputs: list
    ok: bool
    message: str = ""


class ConvertError(Exception):
    pass


class ConvertCancelled(Exception):
    pass


class _Ctx:
    def __init__(self, progress, cancel, index=0, total=1):
        self.progress, self.cancel, self.index, self.total = progress, cancel, index, total
        self.message = ""

    def check(self):
        if self.cancel and self.cancel():
            raise ConvertCancelled()

    def step(self, frac: float, text: str | None = None):
        self.check()
        if self.progress:
            frac = min(max(frac, 0.0), 1.0)
            self.progress(int((self.index + frac) * 1000), self.total * 1000, text)


def unique_path(outdir: str, stem: str, ext: str) -> str:
    stem = re.sub(r'[\\/:*?"<>|]', "_", stem) or "output"
    p = Path(outdir) / f"{stem}.{ext}"
    n = 1
    while p.exists():
        p = Path(outdir) / f"{stem}_{n}.{ext}"
        n += 1
    return str(p)


def _need(module: str, pip_name: str | None = None):
    try:
        return __import__(module)
    except ImportError as ex:
        raise ConvertError(f"缺少组件 {pip_name or module},请先运行:pip install {pip_name or module}") from ex


# ----------------------------------------------------------------------------
# 通用小工具
# ----------------------------------------------------------------------------


def _open_pdf(src: str):
    doc = fitz.open(src)
    if doc.needs_pass:
        doc.close()
        raise ConvertError("这个 PDF 有打开密码。请先在编辑器里打开它(输入密码),再用“添加当前文档”转换。")
    return doc


def _pages(opts: Options, n: int) -> list:
    if not opts.pages.strip():
        return list(range(n))
    seen, out = set(), []
    for g in E.parse_ranges(opts.pages, n):
        for i in g:
            if i not in seen:
                seen.add(i)
                out.append(i)
    return out


def _read_text(path: str) -> str:
    raw = Path(path).read_bytes()
    for enc in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def _save_image(img: Image.Image, path: str, fmt: str, quality: int = 90):
    fmt = fmt.lower()
    has_alpha = img.mode in ("RGBA", "LA") or (img.mode == "P" and "transparency" in img.info)
    if fmt in ("jpg", "jpeg", "bmp"):  # 这两种不支持透明,垫一层白底
        if has_alpha:
            bg = Image.new("RGB", img.size, (255, 255, 255))
            bg.paste(img.convert("RGBA"), mask=img.convert("RGBA").split()[-1])
            img = bg
        else:
            img = img.convert("RGB")
    if fmt in ("jpg", "jpeg"):
        img.save(path, "JPEG", quality=quality, optimize=True)
    elif fmt == "png":
        img.save(path, "PNG", optimize=True)
    elif fmt == "webp":
        img.save(path, "WEBP", quality=quality)
    elif fmt == "bmp":
        img.save(path, "BMP")
    elif fmt == "tiff":
        img.save(path, "TIFF", compression="tiff_lzw")
    elif fmt == "gif":
        img.convert("RGBA" if has_alpha else "RGB").convert("P", palette=Image.ADAPTIVE).save(path, "GIF")
    elif fmt == "ico":
        side = max(img.size)
        square = Image.new("RGBA", (side, side), (0, 0, 0, 0))
        square.paste(img.convert("RGBA"), ((side - img.width) // 2, (side - img.height) // 2))
        sizes = [(s, s) for s in (16, 24, 32, 48, 64, 128, 256) if s <= side] or [(side, side)]
        square.save(path, "ICO", sizes=sizes)
    else:
        raise ConvertError(f"不支持的图片格式:{fmt}")


def _pixmap_image(pm) -> Image.Image:
    return Image.frombytes("RGB", (pm.width, pm.height), pm.samples)


# ----------------------------------------------------------------------------
# OCR(扫描页)
# ----------------------------------------------------------------------------


def _ocr_scanned(doc, idxs, opts: Options, ctx: _Ctx, lo: float, hi: float) -> dict:
    """对扫描页做文字识别,返回 {页号: [{'text','bbox'(pt,显示坐标)}]};没有扫描页或 OCR 不可用时返回 {}。"""
    scanned = [i for i in idxs if E.page_is_scanned(doc[i])]
    if not scanned or not opts.ocr or not ocr.available():
        return {}
    ctx.step(lo, f"正在识别 {len(scanned)} 个扫描页的文字…")
    try:
        res = ocr.ocr_pdf_pages(doc, scanned, opts.ocr_lang, 200,
                                progress=lambda d, t: ctx.step(lo + (hi - lo) * d / t, f"文字识别 {d} / {t} 页…"),
                                cancel=ctx.cancel)
    except ocr.OcrError as ex:
        if "取消" in str(ex):
            raise ConvertCancelled() from ex
        ctx.message += f"(扫描页没能识别文字:{str(ex).splitlines()[0]})"
        return {}
    out = {}
    for pno, (lines, scale) in res.items():
        out[pno] = [{"text": E.fix_punct(ln["text"]), "bbox": [v / scale for v in ln["bbox"]]} for ln in lines]
    return out


# ----------------------------------------------------------------------------
# PDF → 图片 / SVG / 文本 / 网页 / 压缩
# ----------------------------------------------------------------------------


def _pdf_to_images(fmt: str):
    def run(src, outdir, opts, ctx):
        doc = _open_pdf(src)
        idxs = _pages(opts, len(doc))
        stem, outs = Path(src).stem, []
        for k, i in enumerate(idxs):
            ctx.step(k / len(idxs), f"渲染第 {i + 1} 页…")
            img = _pixmap_image(doc[i].get_pixmap(dpi=opts.dpi, alpha=False))
            name = stem if len(idxs) == 1 else f"{stem}_p{i + 1:03d}"
            out = unique_path(outdir, name, "jpg" if fmt == "jpg" else fmt)
            _save_image(img, out, fmt, opts.quality)
            outs.append(out)
        doc.close()
        return outs
    return run


def _pdf_to_svg(src, outdir, opts, ctx):
    doc = _open_pdf(src)
    idxs = _pages(opts, len(doc))
    stem, outs = Path(src).stem, []
    for k, i in enumerate(idxs):
        ctx.step(k / len(idxs), f"导出第 {i + 1} 页…")
        out = unique_path(outdir, stem if len(idxs) == 1 else f"{stem}_p{i + 1:03d}", "svg")
        Path(out).write_text(doc[i].get_svg_image(text_as_path=True), encoding="utf-8")
        outs.append(out)
    doc.close()
    return outs


def _pdf_to_txt(src, outdir, opts, ctx):
    doc = _open_pdf(src)
    idxs = _pages(opts, len(doc))
    ocr_pages = _ocr_scanned(doc, idxs, opts, ctx, 0.0, 0.7)
    parts = []
    for k, i in enumerate(idxs):
        ctx.step(0.7 + 0.3 * k / len(idxs), f"提取第 {i + 1} 页文字…")
        if i in ocr_pages:
            text = "\n".join(ln["text"] for ln in ocr_pages[i])
        else:
            text = doc[i].get_text(sort=True).replace(" ", " ")
        parts.append(text.rstrip())
    out = unique_path(outdir, Path(src).stem, "txt")
    Path(out).write_text(("\n\n" if len(parts) > 1 else "").join(
        (f"—— 第 {i + 1} 页 ——\n{t}" if len(parts) > 1 else t) for i, t in zip(idxs, parts)), encoding="utf-8")
    doc.close()
    return [out]


def _pdf_to_html(src, outdir, opts, ctx):
    doc = _open_pdf(src)
    idxs = _pages(opts, len(doc))
    ocr_pages = _ocr_scanned(doc, idxs, opts, ctx, 0.0, 0.6)
    parts = []
    for k, i in enumerate(idxs):
        ctx.step(0.6 + 0.4 * k / len(idxs), f"转换第 {i + 1} 页…")
        if i in ocr_pages:
            parts.append("".join(f"<p>{htmllib.escape(ln['text'])}</p>" for ln in ocr_pages[i]))
        else:
            parts.append(doc[i].get_text("xhtml"))
    title = htmllib.escape(Path(src).stem)
    body = '\n<hr class="pagebreak">\n'.join(parts)
    # PyMuPDF 把中文写成 &#x793a; 这样的数字实体;还原成真正的字符,方便直接阅读和编辑源码(&lt; &amp; 保持不变)
    body = re.sub(r"&#(x[0-9a-fA-F]+|\d+);",
                  lambda m: chr(int(m.group(1)[1:], 16) if m.group(1)[0] == "x" else int(m.group(1)))
                  if (int(m.group(1)[1:], 16) if m.group(1)[0] == "x" else int(m.group(1))) > 127 else m.group(0), body)
    out = unique_path(outdir, Path(src).stem, "html")
    Path(out).write_text(
        f'<!DOCTYPE html>\n<html lang="zh"><head><meta charset="utf-8"><title>{title}</title><style>'
        'body{font-family:"Microsoft YaHei",Arial,sans-serif;max-width:920px;margin:2em auto;line-height:1.7;padding:0 1em}'
        "img{max-width:100%}hr.pagebreak{margin:2.5em 0;border:0;border-top:1px dashed #bbb}</style></head>"
        f"<body>\n{body}\n</body></html>", encoding="utf-8")
    doc.close()
    return [out]


def _pdf_compress(src, outdir, opts, ctx):
    doc = _open_pdf(src)
    ctx.step(0.1, "正在压缩图片…")
    try:
        doc.rewrite_images(dpi_threshold=opts.dpi + 30, dpi_target=opts.dpi, quality=opts.quality)
    except Exception:  # noqa: BLE001  旧版本 PyMuPDF 没有这个功能
        pass
    ctx.step(0.6, "正在优化字体和结构…")
    E.subset_fonts(doc)
    out = unique_path(outdir, Path(src).stem + "_压缩", "pdf")
    doc.save(out, garbage=4, deflate=True, clean=True)
    doc.close()
    before, after = os.path.getsize(src), os.path.getsize(out)
    ctx.message = f"{before / 1024:.0f} KB → {after / 1024:.0f} KB" + ("" if after < before else "(没有变小,原文件可能已经很精简)")
    return [out]


# ----------------------------------------------------------------------------
# PDF → Word(版面分析和写出在 pdf_to_word.py)
# ----------------------------------------------------------------------------


def _pdf_to_docx(src, outdir, opts, ctx):
    _need("docx", "python-docx")
    import pdf_to_word

    doc = _open_pdf(src)
    idxs = _pages(opts, len(doc))
    ocr_pages = _ocr_scanned(doc, idxs, opts, ctx, 0.0, 0.4)
    d, warnings = pdf_to_word.build_document(doc, idxs, ocr_pages, lambda f, t=None: ctx.step(0.4 + 0.6 * f, t))
    out = unique_path(outdir, Path(src).stem, "docx")
    d.save(out)
    doc.close()
    if warnings:
        ctx.message = ";".join(warnings)
    return [out]


# ----------------------------------------------------------------------------
# PDF → Excel / PowerPoint
# ----------------------------------------------------------------------------

_NUM = re.compile(r"^-?\d+(\.\d+)?$|^-?\d{1,3}(,\d{3})+(\.\d+)?$")


def _cell_value(v):
    if v is None:
        return ""
    s = str(v).strip().replace(" ", " ")
    if _NUM.match(s) and not (len(s) > 1 and s.startswith("0") and not s.startswith("0.")):
        try:
            f = float(s.replace(",", ""))
            return int(f) if f.is_integer() and "." not in s else f
        except ValueError:
            pass
    return s


def _sheet_title(wb, title: str) -> str:
    title = re.sub(r"[\[\]:*?/\\]", "_", title)[:31]
    base, n = title, 2
    while title in wb.sheetnames:
        title = f"{base[:28]}_{n}"
        n += 1
    return title


def _pdf_to_xlsx(src, outdir, opts, ctx):
    _need("openpyxl")
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    doc = _open_pdf(src)
    idxs = _pages(opts, len(doc))
    wb = Workbook()
    wb.remove(wb.active)
    n_tables = 0
    for k, i in enumerate(idxs):
        ctx.step(0.7 * k / len(idxs), f"查找第 {i + 1} 页的表格…")
        try:
            tabs = doc[i].find_tables().tables
        except Exception:  # noqa: BLE001
            tabs = []
        for ti, t in enumerate(tabs):
            rows = t.extract()
            if not rows:
                continue
            ws = wb.create_sheet(_sheet_title(wb, f"第{i + 1}页-表{ti + 1}"))
            for row in rows:
                ws.append([_cell_value(c) for c in row])
            for ci in range(1, max(len(r) for r in rows) + 1):
                width = max((len(str(ws.cell(r, ci).value or "")) for r in range(1, len(rows) + 1)), default=8)
                ws.column_dimensions[get_column_letter(ci)].width = min(max(width * 1.6 + 2, 8), 60)
            for r in ws.iter_rows():
                for c in r:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
            n_tables += 1
    if n_tables == 0:
        ocr_pages = _ocr_scanned(doc, idxs, opts, ctx, 0.7, 0.95)
        ws = wb.create_sheet("文本")
        for i in idxs:
            lines = ([ln["text"] for ln in ocr_pages[i]] if i in ocr_pages
                     else doc[i].get_text(sort=True).replace(" ", " ").splitlines())
            for line in lines:
                ws.append([_cell_value(line) if _NUM.match(line.strip()) else line])
        ws.column_dimensions["A"].width = 100
        ctx.message = "没有检测到表格,已把文字按行放进工作表"
    else:
        ctx.message = f"共找到 {n_tables} 个表格"
    out = unique_path(outdir, Path(src).stem, "xlsx")
    wb.save(out)
    doc.close()
    return [out]


def _jpeg_bytes(img: Image.Image, quality: int) -> bytes:
    buf = io.BytesIO()
    img.convert("RGB").save(buf, "JPEG", quality=quality)
    return buf.getvalue()


def _render_without_text(doc, pno: int, dpi: int) -> Image.Image:
    """渲染去掉全部文字后的页面(保留图片和图形),作为幻灯片背景。"""
    tmp = fitz.open()
    tmp.insert_pdf(doc, from_page=pno, to_page=pno)
    E._apply_redactions(tmp[0], [list(tmp[0].rect)], fill=False, remove_graphics=False)
    img = _pixmap_image(tmp[0].get_pixmap(dpi=dpi, alpha=False))
    tmp.close()
    return img


def _pdf_to_pptx(src, outdir, opts, ctx):
    _need("pptx", "python-pptx")
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt

    import pdf_to_word

    doc = _open_pdf(src)
    idxs = _pages(opts, len(doc))
    fonts = pdf_to_word.Fonts(doc)  # 用 PDF 里字体的真实名字
    prs = Presentation()
    W0, H0 = doc[idxs[0]].rect.width, doc[idxs[0]].rect.height
    k0 = min(1.0, 4000 / max(W0, H0))
    prs.slide_width, prs.slide_height = Emu(int(W0 * k0 * 12700)), Emu(int(H0 * k0 * 12700))
    blank = prs.slide_layouts[6]
    editable = opts.pptx_mode == "editable"
    for k, i in enumerate(idxs):
        ctx.step(k / len(idxs), f"转换第 {i + 1} 页…")
        page = doc[i]
        slide = prs.slides.add_slide(blank)
        W, H = page.rect.width, page.rect.height
        fit = min(W0 * k0 / W, H0 * k0 / H)  # 页面尺寸不同时缩放进幻灯片
        offx, offy = (W0 * k0 - W * fit) / 2, (H0 * k0 - H * fit) / 2
        emu = lambda v: Emu(int(v * 12700))  # noqa: E731
        has_text = E.page_has_text(page)
        if editable and has_text:
            bg = _render_without_text(doc, i, opts.dpi)
        else:
            bg = _pixmap_image(page.get_pixmap(dpi=opts.dpi, alpha=False))
        slide.shapes.add_picture(io.BytesIO(_jpeg_bytes(bg, opts.quality)), emu(offx), emu(offy),
                                 emu(W * fit), emu(H * fit))
        text_all = page.get_text().replace(" ", " ")
        if editable and has_text:
            rm = page.rotation_matrix
            for b in page.get_text("dict", flags=E._TEXT_FLAGS)["blocks"]:
                if b["type"] != 0:
                    continue
                for ln in b["lines"]:
                    spans = [s for s in ln["spans"] if s["text"]]
                    if not spans or not "".join(s["text"] for s in spans).strip():
                        continue
                    r = (fitz.Rect(ln["bbox"]) * rm).normalize()
                    tb = slide.shapes.add_textbox(emu(offx + r.x0 * fit), emu(offy + r.y0 * fit),
                                                  emu(r.width * fit * 1.08 + 6), emu(r.height * fit))
                    tf = tb.text_frame
                    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                    tf.word_wrap = False
                    para = tf.paragraphs[0]
                    for s in spans:
                        run = para.add_run()
                        run.text = s["text"].replace(" ", " ")
                        label, bold, italic = fonts.info(page, s)
                        run.font.size = Pt(max(s["size"] * fit, 1))
                        run.font.name = label
                        run.font.bold = bold or None
                        run.font.italic = italic or None
                        run.font.color.rgb = RGBColor((s["color"] >> 16) & 255, (s["color"] >> 8) & 255, s["color"] & 255)
        if text_all.strip():
            slide.notes_slide.notes_text_frame.text = text_all.strip()
    out = unique_path(outdir, Path(src).stem, "pptx")
    prs.save(out)
    doc.close()
    ctx.message = "文字为可编辑文本框" if editable else "每页是一张图片"
    return [out]


# ----------------------------------------------------------------------------
# HTML / 文本 → PDF
# ----------------------------------------------------------------------------

_BASE_CSS = ("body{font-size:11pt;line-height:1.55} h1{font-size:20pt} h2{font-size:16pt} h3{font-size:13pt}"
             "pre,code{font-family:monospace;font-size:10pt;background:#f3f3f3} table{border-collapse:collapse}"
             "td,th{border:1px solid #888;padding:3pt 5pt} blockquote{margin-left:12pt;color:#555}"
             "img{max-width:100%} p{margin:0 0 6pt 0}")


def _font_css() -> tuple:
    """(css, 字体目录):用系统的微软雅黑,中文才能正常显示和嵌入。"""
    reg = E._find_font_file("msyh.ttc") or E._find_font_file("simhei.ttf") or E._find_font_file("simsun.ttc")
    if not reg:
        return "", None
    name = os.path.basename(reg)
    css = f'@font-face{{font-family:"CJK";src:url({name})}}'
    bold = E._find_font_file("msyhbd.ttc") if name == "msyh.ttc" else None
    if bold and os.path.dirname(bold) == os.path.dirname(reg):
        css += f'@font-face{{font-family:"CJK";font-weight:bold;src:url({os.path.basename(bold)})}}'
    return css + 'body,p,li,td,th,h1,h2,h3,h4,h5,h6{font-family:"CJK",sans-serif}', os.path.dirname(reg)


def _html_to_pdf(html_text: str, out: str, base_dir: str | None = None, landscape: bool = False, ctx=None):
    css, fontdir = _font_css()
    arch = fitz.Archive()
    if fontdir:
        arch.add(fitz.Archive(fontdir))
    if base_dir and os.path.isdir(base_dir):
        arch.add(fitz.Archive(base_dir))
    story = fitz.Story(html=html_text, user_css=css + _BASE_CSS, archive=arch)
    media = fitz.paper_rect("a4-l" if landscape else "a4")
    where = media + (48, 48, -48, -48)
    buf = io.BytesIO()  # 写进内存,避免 Windows 下临时文件被占用
    writer = fitz.DocumentWriter(buf)
    more, pages = 1, 0
    while more:
        if ctx:
            ctx.check()
        dev = writer.begin_page(media)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
        pages += 1
        if pages > 3000:
            break
    writer.close()
    d = fitz.open(stream=buf.getvalue(), filetype="pdf")
    E.subset_fonts(d)  # 完整的微软雅黑有十几 MB,裁剪后只剩用到的字
    d.save(out, garbage=3, deflate=True)
    d.close()


def _text_to_html(text: str) -> str:
    out = []
    for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        stripped = line.lstrip(" ")
        lead = "&nbsp;" * (len(line) - len(stripped))
        out.append(f"<p>{lead}{htmllib.escape(stripped.replace(chr(9), '    ')) or '&nbsp;'}</p>")
    return "".join(out)


def _text_to_pdf(src, outdir, opts, ctx):
    ctx.step(0.1, "正在排版…")
    ext = Path(src).suffix.lower()
    text = _read_text(src)
    base = os.path.dirname(os.path.abspath(src))
    if ext in (".html", ".htm"):
        body = text
    elif ext in (".md", ".markdown"):
        try:
            import markdown
            body = markdown.markdown(text, extensions=["tables", "fenced_code", "sane_lists"])
        except ImportError:
            body = _text_to_html(text)
    else:
        body = _text_to_html(text)
    out = unique_path(outdir, Path(src).stem, "pdf")
    _html_to_pdf(body, out, base, ctx=ctx)
    return [out]


# ----------------------------------------------------------------------------
# Office → PDF
# ----------------------------------------------------------------------------

_PS_OFFICE = r"""
param([string]$In, [string]$Out, [string]$Kind)
$ErrorActionPreference = 'Stop'
[Console]::OutputEncoding = New-Object System.Text.UTF8Encoding($false)
function Pick($ids) { foreach ($i in $ids) { if (Test-Path "Registry::HKEY_CLASSES_ROOT\$i") { return $i } } return $null }
$ids = @{ word = @('Word.Application','KWPS.Application','wps.Application')
          excel = @('Excel.Application','KET.Application','et.Application')
          ppt = @('PowerPoint.Application','KWPP.Application','wpp.Application') }
$prog = Pick $ids[$Kind]
if (-not $prog) { Write-Output 'NOAPP'; exit 3 }
$app = New-Object -ComObject $prog
$code = 0
try {
    if ($Kind -eq 'word') {
        $app.Visible = $false
        try { $app.DisplayAlerts = 0 } catch {}
        $doc = $app.Documents.Open($In, $false, $true)
        try { $doc.ExportAsFixedFormat($Out, 17) } catch { $doc.SaveAs2($Out, 17) }
        $doc.Close($false)
    } elseif ($Kind -eq 'excel') {
        $app.Visible = $false
        $app.DisplayAlerts = $false
        $wb = $app.Workbooks.Open($In, 0, $true)
        foreach ($ws in $wb.Worksheets) {
            try { $ws.PageSetup.Zoom = $false; $ws.PageSetup.FitToPagesWide = 1; $ws.PageSetup.FitToPagesTall = $false } catch {}
        }
        $wb.ExportAsFixedFormat(0, $Out)
        $wb.Close($false)
    } else {
        $pres = $app.Presentations.Open($In, -1, 0, 0)
        $pres.SaveAs($Out, 32)
        $pres.Close()
    }
    Write-Output 'OK'
} catch {
    Write-Output ('ERR ' + $_.Exception.Message)
    $code = 4
} finally {
    try { $app.Quit() } catch {}
    try { [System.Runtime.InteropServices.Marshal]::ReleaseComObject($app) | Out-Null } catch {}
}
exit $code
"""
_NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0)


def _run_process(cmd: list, timeout: int, ctx: _Ctx):
    """运行外部程序,可取消、有超时。返回 (退出码, 输出文本)。"""
    proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, creationflags=_NO_WINDOW)
    import time
    t0 = time.time()
    while True:
        try:
            out, _ = proc.communicate(timeout=0.4)
            return proc.returncode, out.decode("utf-8", "replace")
        except subprocess.TimeoutExpired:
            if (ctx.cancel and ctx.cancel()) or time.time() - t0 > timeout:
                proc.kill()
                proc.communicate()
                if time.time() - t0 > timeout:
                    raise ConvertError("转换超时(超过 %d 秒)" % timeout)
                raise ConvertCancelled()


def _office_com(kind: str, src: str, out: str, ctx: _Ctx):
    if sys.platform != "win32":
        raise ConvertError("只能在 Windows 上通过 Office / WPS 转换")
    # Office 对太长的路径、中文路径很挑剔:先复制到临时目录,用短的英文名转换,再把结果移回去
    work = tempfile.mkdtemp(prefix="pdfed_")
    tmp_in = os.path.join(work, "in" + Path(src).suffix.lower())
    tmp_out = os.path.join(work, "out.pdf")
    shutil.copyfile(src, tmp_in)
    script = os.path.join(work, "run.ps1")
    with open(script, "w", encoding="utf-8-sig") as f:
        f.write(_PS_OFFICE)
    try:
        code, text = _run_process(["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File",
                                   script, "-In", tmp_in, "-Out", tmp_out, "-Kind", kind], 300, ctx)
        text = text.strip()
        if "NOAPP" in text:
            raise ConvertError("没有找到 Microsoft Office / WPS")
        if not os.path.exists(tmp_out):
            raise ConvertError(text.replace("ERR ", "")[:300] or "Office 没有生成 PDF")
        shutil.move(tmp_out, out)
    finally:
        shutil.rmtree(work, ignore_errors=True)


def _find_soffice():
    p = shutil.which("soffice")
    if p:
        return p
    for base in (os.environ.get("ProgramFiles", ""), os.environ.get("ProgramFiles(x86)", "")):
        cand = os.path.join(base, "LibreOffice", "program", "soffice.exe")
        if base and os.path.exists(cand):
            return cand
    return None


def _libreoffice(src: str, out: str, ctx: _Ctx):
    so = _find_soffice()
    if not so:
        raise ConvertError("没有找到 LibreOffice")
    tmp = tempfile.mkdtemp(prefix="pdfed_lo_")
    try:
        tmp_in = os.path.join(tmp, "in" + Path(src).suffix.lower())
        shutil.copyfile(src, tmp_in)
        code, text = _run_process([so, "--headless", "--norestore", "--convert-to", "pdf", "--outdir", tmp, tmp_in],
                                  300, ctx)
        produced = os.path.join(tmp, "in.pdf")
        if not os.path.exists(produced):
            raise ConvertError("LibreOffice 转换失败:" + text.strip()[:200])
        shutil.move(produced, out)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _docx_builtin(src: str, out: str, ctx: _Ctx):
    if Path(src).suffix.lower() != ".docx":
        raise ConvertError("内置方式只支持 .docx")
    try:
        import mammoth  # 显式 import:PyInstaller 才能发现并打包它
    except ImportError as ex:
        raise ConvertError("缺少组件 mammoth,请先运行:pip install mammoth") from ex
    with open(src, "rb") as f:
        result = mammoth.convert_to_html(f)
    _html_to_pdf(f"<html><body>{result.value}</body></html>", out, ctx=ctx)


def _sheet_builtin(src: str, out: str, ctx: _Ctx):
    ext = Path(src).suffix.lower()
    if ext == ".csv":
        import csv
        text = _read_text(src)
        rows = list(csv.reader(io.StringIO(text)))
        sheets = [(Path(src).stem, rows)]
    elif ext == ".xlsx":
        _need("openpyxl")
        import openpyxl
        wb = openpyxl.load_workbook(src, data_only=True)
        sheets = [(ws.title, [["" if c is None else c for c in row] for row in ws.iter_rows(values_only=True)])
                  for ws in wb.worksheets]
    else:
        raise ConvertError("内置方式只支持 .xlsx / .csv")
    parts = []
    for name, rows in sheets:
        body = "".join("<tr>" + "".join(f"<td>{htmllib.escape(str(c))}</td>" for c in r) + "</tr>" for r in rows)
        parts.append(f"<h3>{htmllib.escape(name)}</h3><table>{body}</table>")
    _html_to_pdf("<html><body>" + "".join(parts) + "</body></html>", out, landscape=True, ctx=ctx)


def office_engines() -> list:
    """本机能把 Word/Excel/PPT 转成 PDF 的软件,用于在界面上提示用户。"""
    found = []
    if sys.platform == "win32":
        import winreg

        def has(progid):
            try:
                winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, progid))
                return True
            except OSError:
                return False
        if any(has(p) for p in ("Word.Application", "Excel.Application", "PowerPoint.Application")):
            found.append("Microsoft Office")
        if any(has(p) for p in ("KWPS.Application", "Kwps.Application", "KET.Application", "KWPP.Application")):
            found.append("WPS Office")
    if _find_soffice():
        found.append("LibreOffice")
    return found


def _office_to_pdf_file(group: str, src: str, out: str, ctx: _Ctx) -> str:
    """把 Office 文件转成 PDF 写到 out。依次尝试:Office/WPS → LibreOffice → 内置。返回使用的方式名。"""
    kind = {"word": "word", "excel": "excel", "ppt": "ppt"}[group]
    errors = []
    is_csv = Path(src).suffix.lower() == ".csv"
    if not is_csv:
        for name, fn in (("Office/WPS", lambda: _office_com(kind, src, out, ctx)),
                         ("LibreOffice", lambda: _libreoffice(src, out, ctx))):
            ctx.step(0.15, f"正在用 {name} 转换…")
            try:
                fn()
                ctx.message = f"使用 {name} 转换"
                return name
            except ConvertError as ex:
                if "没有找到" not in str(ex):  # 没装的软件不算"失败原因",不打扰用户
                    errors.append(f"{name} 失败:{str(ex).strip().splitlines()[0][:120]}")
                if os.path.exists(out):
                    os.unlink(out)
    ctx.step(0.4, "正在用内置方式转换…")
    try:
        if group == "word":
            _docx_builtin(src, out, ctx)
        elif group == "excel":
            _sheet_builtin(src, out, ctx)
        else:
            raise ConvertError("PowerPoint 文件需要安装 Microsoft Office、WPS 或 LibreOffice 才能转换")
        ctx.message = "使用内置方式转换(版式比 Office 简单)" + (";" + ";".join(errors) if errors else "")
        return "内置"
    except ConvertError as ex:
        if not errors:
            errors.append("没有找到 Microsoft Office / WPS / LibreOffice")
        errors.append(f"内置:{ex}")
    raise ConvertError("转换失败。\n" + "\n".join(errors))


def _office_to_pdf(group: str):
    def run(src, outdir, opts, ctx):
        out = unique_path(outdir, Path(src).stem, "pdf")
        _office_to_pdf_file(group, src, out, ctx)
        return [out]
    return run


def _office_to_images(group: str, fmt: str):
    def run(src, outdir, opts, ctx):
        tmp = tempfile.mkdtemp(prefix="pdfeditor_off_")
        try:
            pdf = os.path.join(tmp, Path(src).stem + ".pdf")
            _office_to_pdf_file(group, src, pdf, ctx)
            return _pdf_to_images(fmt)(pdf, outdir, opts, ctx)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
    return run


def _docx_to_txt(src, outdir, opts, ctx):
    if Path(src).suffix.lower() != ".docx":
        raise ConvertError("转纯文本目前只支持 .docx(.doc 请先另存为 .docx)")
    _need("docx", "python-docx")
    from docx import Document
    d = Document(src)
    lines = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            lines.append("\t".join(c.text for c in row.cells))
    out = unique_path(outdir, Path(src).stem, "txt")
    Path(out).write_text("\n".join(lines), encoding="utf-8")
    return [out]


# ----------------------------------------------------------------------------
# 图片 → PDF / 其它图片格式
# ----------------------------------------------------------------------------


def _load_image(path: str):
    img = Image.open(path)
    img.load()
    rotated = img.getexif().get(274, 1) not in (1, None)
    return ImageOps.exif_transpose(img), rotated


def _add_image_page(doc, path: str, opts: Options):
    img, rotated = _load_image(path)
    w, h = img.size
    if opts.page_size == "fit":
        dpi = img.info.get("dpi", (opts.dpi, opts.dpi))[0] or opts.dpi
        pw, ph = w * 72 / max(dpi, 36), h * 72 / max(dpi, 36)
        rect = fitz.Rect(0, 0, pw, ph)
    else:
        pw, ph = (841.89, 595.28) if w > h else (595.28, 841.89)
        m = 24
        k = min((pw - 2 * m) / w, (ph - 2 * m) / h)
        iw, ih = w * k, h * k
        rect = fitz.Rect((pw - iw) / 2, (ph - ih) / 2, (pw + iw) / 2, (ph + ih) / 2)
    page = doc.new_page(width=pw, height=ph)
    ext = Path(path).suffix.lower()
    if ext in (".jpg", ".jpeg", ".png") and not rotated and img.mode in ("RGB", "L", "RGBA", "P", "LA"):
        page.insert_image(rect, filename=path)  # 原样嵌入,不重新压缩
    else:
        buf = io.BytesIO()
        if img.mode in ("RGBA", "LA", "P"):
            img.convert("RGBA").save(buf, "PNG")
        else:
            img.convert("RGB").save(buf, "JPEG", quality=opts.quality)
        page.insert_image(rect, stream=buf.getvalue())


def _images_to_pdf_doc(paths: list, opts: Options, ctx: _Ctx):
    doc = fitz.open()
    for k, p in enumerate(paths):
        ctx.step(k / max(len(paths), 1), f"添加 {Path(p).name}…")
        _add_image_page(doc, p, opts)
    return doc


def _image_to_pdf(src, outdir, opts, ctx):
    doc = _images_to_pdf_doc([src], opts, ctx)
    out = unique_path(outdir, Path(src).stem, "pdf")
    doc.save(out, garbage=3, deflate=True)
    doc.close()
    return [out]


def merge_images_to_pdf(paths: list, outdir: str, opts: Options, progress=None, cancel=None) -> Result:
    ctx = _Ctx(progress, cancel)
    try:
        os.makedirs(outdir, exist_ok=True)
        doc = _images_to_pdf_doc(paths, opts, ctx)
        out = unique_path(outdir, f"{Path(paths[0]).stem}_等{len(paths)}张图片", "pdf")
        doc.save(out, garbage=3, deflate=True)
        doc.close()
        return Result(paths[0], [out], True, f"{len(paths)} 张图片合并成 1 个 PDF")
    except ConvertCancelled:
        raise
    except Exception as ex:  # noqa: BLE001
        return Result(paths[0], [], False, str(ex))


def _image_to_image(fmt: str):
    def run(src, outdir, opts, ctx):
        ctx.step(0.3, "正在转换…")
        img, _ = _load_image(src)
        out = unique_path(outdir, Path(src).stem, "jpg" if fmt == "jpg" else fmt)
        _save_image(img, out, fmt, opts.quality)
        return [out]
    return run


# ----------------------------------------------------------------------------
# 调度
# ----------------------------------------------------------------------------


def _handler(group: str, target: str):
    if group == "pdf":
        if target == "docx":
            return _pdf_to_docx
        if target == "xlsx":
            return _pdf_to_xlsx
        if target == "pptx":
            return _pdf_to_pptx
        if target in ("png", "jpg", "webp", "tiff", "bmp"):
            return _pdf_to_images(target)
        if target == "txt":
            return _pdf_to_txt
        if target == "html":
            return _pdf_to_html
        if target == "svg":
            return _pdf_to_svg
        if target == "pdf":
            return _pdf_compress
    elif group in ("word", "excel", "ppt"):
        if target == "pdf":
            return _office_to_pdf(group)
        if target in ("png", "jpg"):
            return _office_to_images(group, target)
        if target == "txt" and group == "word":
            return _docx_to_txt
    elif group == "image":
        if target == "pdf":
            return _image_to_pdf
        if target in IMAGE_FORMATS:
            return _image_to_image(target)
    elif group == "text" and target == "pdf":
        return _text_to_pdf
    raise ConvertError(f"不支持把 {GROUP_NAMES.get(group, group)} 转成 {target}")


def convert(src: str, target: str, outdir: str, opts: Options | None = None, progress=None, cancel=None,
            index: int = 0, total: int = 1) -> Result:
    """转换一个文件。出错不抛异常,返回 ok=False 的 Result(用户取消会抛 ConvertCancelled)。"""
    opts = opts or Options()
    ctx = _Ctx(progress, cancel, index, total)
    group = group_of(src)
    try:
        if group is None:
            raise ConvertError("不支持这种文件类型")
        if not os.path.exists(src):
            raise ConvertError("文件不存在")
        os.makedirs(outdir, exist_ok=True)
        ctx.step(0.0, f"正在转换 {Path(src).name}…")
        outs = _handler(group, target)(src, outdir, opts, ctx)
        ctx.step(1.0)
        return Result(src, outs, True, ctx.message)
    except ConvertCancelled:
        raise
    except ConvertError as ex:
        return Result(src, [], False, str(ex))
    except Exception as ex:  # noqa: BLE001
        return Result(src, [], False, f"{type(ex).__name__}: {ex}")


def convert_many(srcs: list, target: str, outdir_for, opts: Options | None = None, progress=None, cancel=None) -> list:
    """批量转换。outdir_for(src) 给出每个文件的输出目录。用户取消时返回已完成的结果。"""
    opts = opts or Options()
    results = []
    try:
        if len(srcs) > 1 and target == "pdf" and opts.merge and all(group_of(s) == "image" for s in srcs):
            return [merge_images_to_pdf(srcs, outdir_for(srcs[0]), opts, progress, cancel)]
        for i, s in enumerate(srcs):
            results.append(convert(s, target, outdir_for(s), opts, progress, cancel, i, len(srcs)))
    except ConvertCancelled:
        results.append(Result("", [], False, "已取消"))
    return results


# ----------------------------------------------------------------------------
# 命令行:python main.py --convert 文件 --to docx [--out 目录] [--pages 1-3] [--dpi 150]
# ----------------------------------------------------------------------------


def cli(argv: list) -> int:
    import argparse
    ap = argparse.ArgumentParser(prog="convert", description="文件格式转换")
    ap.add_argument("files", nargs="+")
    ap.add_argument("--to", required=True, help="目标格式:" + ",".join(sorted({k for t in TARGETS.values() for k, _ in t})))
    ap.add_argument("--out", help="输出目录(默认与源文件相同)")
    ap.add_argument("--pages", default="")
    ap.add_argument("--dpi", type=int, default=150)
    ap.add_argument("--quality", type=int, default=90)
    ap.add_argument("--merge", action="store_true", help="多张图片合并成一个 PDF")
    ap.add_argument("--no-ocr", action="store_true")
    args = ap.parse_args(argv)
    opts = Options(pages=args.pages, dpi=args.dpi, quality=args.quality, merge=args.merge, ocr=not args.no_ocr)
    results = convert_many(args.files, args.to, lambda s: args.out or os.path.dirname(os.path.abspath(s)), opts)
    bad = 0
    for r in results:
        if r.ok:
            print("OK  ", r.src, "->", ", ".join(r.outputs), r.message)
        else:
            bad += 1
            print("FAIL", r.src, r.message)
    return 1 if bad else 0
