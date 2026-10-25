"""PDF → Word(.docx):版面分析与写出。

做法:先把每一页分析成"元素"——段落、表格、图片;再按阅读顺序写进 Word。
要点:
· 字体:读取 PDF 内嵌字体的真实字体名(微软雅黑、DengXian、Cambria……),不统一替换成别的字体;
· 项目符号 / 编号:把 "•"、"1." 与后面的文字并成同一行,写成带悬挂缩进和制表位的列表段落;
  矢量画出来的小圆点也识别成项目符号;Symbol / Wingdings 的私有字符还原成 "•";
· 图片:JPEG 分辨率字段为 0 会让 python-docx 报错(以前被悄悄吞掉了),这里修正;
  倒置/翻转/带遮罩的图片、PDF 里没有编号的图片,改用"去掉文字后的页面渲染"截取;
  与文字重叠的图和大背景图做成"衬于文字下方"的浮动图片;
· 矢量图(图表、卡片、色块):整块渲染成图片;里面文字不多时连文字一起渲染,否则做成背景图、文字保持可编辑;
· 扫描页 / 带不可见文字层的扫描件:用识别出的文字重新排成段落。
"""
from __future__ import annotations

import io
import re
from collections import Counter

import engine as E
from PIL import Image

fitz = E.fitz

# ----------------------------------------------------------------------------
# 字体名
# ----------------------------------------------------------------------------

_PS_FAMILY = {
    "arial": "Arial", "arialnarrow": "Arial Narrow", "arialblack": "Arial Black", "helvetica": "Arial",
    "timesnewroman": "Times New Roman", "times": "Times New Roman", "timesroman": "Times New Roman",
    "couriernew": "Courier New", "courier": "Courier New", "symbol": "Symbol", "wingdings": "Wingdings",
    "calibri": "Calibri", "cambria": "Cambria", "verdana": "Verdana", "tahoma": "Tahoma", "georgia": "Georgia",
    "consolas": "Consolas", "segoeui": "Segoe UI", "trebuchetms": "Trebuchet MS", "garamond": "Garamond",
    "microsoftyahei": "Microsoft YaHei", "microsoftyaheiui": "Microsoft YaHei UI", "msyh": "Microsoft YaHei",
    "simsun": "SimSun", "nsimsun": "NSimSun", "宋体": "SimSun", "simhei": "SimHei", "黑体": "SimHei",
    "kaiti": "KaiTi", "simkai": "KaiTi", "楷体": "KaiTi", "fangsong": "FangSong", "simfang": "FangSong",
    "仿宋": "FangSong", "dengxian": "DengXian", "等线": "DengXian", "msmincho": "MS Mincho", "msgothic": "MS Gothic",
    "stsong": "STSong", "stheiti": "STHeiti", "stkaiti": "STKaiti", "stfangsong": "STFangsong",
    "pingfangsc": "PingFang SC", "sourcehansanssc": "Source Han Sans SC", "notosanssc": "Noto Sans SC",
    "droidsansfallback": "Microsoft YaHei",
}
_STYLE_WORDS = ("bolditalic", "boldoblique", "semibold", "demibold", "extrabold", "bold", "italic", "oblique",
                "regular", "light", "medium", "heavy", "thin")


def _compact(name: str) -> str:
    s = re.sub(r"^[A-Z]{6}\+", "", name or "")
    return re.sub(r"[\s\-_,]", "", s).lower()


def _strip_style(c: str) -> str:
    for w in _STYLE_WORDS:
        c = c.replace(w, "")
    return re.sub(r"(psmt|mt|ps)$", "", c)


class Fonts:
    """把 PDF 里的字体(PostScript 名)翻译成 Word 认识的字体名。优先读内嵌字体文件里的真实名字。"""

    def __init__(self, doc):
        self.doc = doc
        self._cache: dict = {}
        self._xrefs: dict = {}

    def _xref_of(self, page, basefont: str):
        if page.number not in self._xrefs:
            try:
                self._xrefs[page.number] = {f[3]: f[0] for f in page.get_fonts(full=True)}
            except Exception:  # noqa: BLE001
                self._xrefs[page.number] = {}
        return self._xrefs[page.number].get(basefont)

    def _embedded_name(self, xref):
        try:
            _name, ext, _typ, buf = self.doc.extract_font(xref)
            if not buf or ext not in ("ttf", "otf", "ttc"):
                return None
            from fontTools.ttLib import TTFont
            font = TTFont(io.BytesIO(buf), lazy=True)
            recs = font["name"].names
            for nid in (16, 1):
                cands = [r for r in recs if r.nameID == nid]
                cands.sort(key=lambda r: (not (r.platformID == 3 and r.langID == 0x409), r.platformID != 1))
                for r in cands:
                    txt = r.toUnicode().strip()
                    if txt and not txt.lower().startswith(("untitled", "unknown")):
                        return txt
        except Exception:  # noqa: BLE001
            return None
        return None

    def family(self, page, basefont: str) -> str:
        key = basefont
        if key in self._cache:
            return self._cache[key]
        fam = None
        comp = _strip_style(_compact(basefont))
        if comp in _PS_FAMILY:
            fam = _PS_FAMILY[comp]
        if fam is None:
            xref = self._xref_of(page, basefont)
            if xref:
                fam = self._embedded_name(xref)
        if fam is None:
            clean = re.sub(r"^[A-Z]{6}\+", "", basefont or "")
            clean = re.split(r"[-,]", clean)[0]
            clean = re.sub(r"(PSMT|MT|PS)$", "", clean)
            fam = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", clean).strip() or "Arial"
        self._cache[key] = fam
        return fam

    def info(self, page, span):
        """(Word 字体名, 粗体, 斜体)"""
        _k, bold, italic = E.guess_font_style(span["font"], span["flags"])
        return self.family(page, span["font"]), bold, italic


# ----------------------------------------------------------------------------
# 文字清理 / 项目符号
# ----------------------------------------------------------------------------

_PRIVATE = {"": "•", "": "▪", "": "➢", "": "❖", "": "✓", "": "●", "": "■",
            "": "☐", "": "➔", "": "✗", "": "◆", "": "○", "": "•", "": "•"}
_BULLETS = set("•·●○◦▪■□◆◇▶►➢➤✓✔❖★☆–—-*➔☐✗")
_NUM_MARK = re.compile(r"^(\(?\d{1,3}[.)、）]|[a-zA-Z][.)]|[一二三四五六七八九十]+[、.]|[①-⑩]|\(\d+\)|[ivxIVX]+[.)])$")


def clean_text(text: str) -> str:
    text = text.replace(" ", " ").replace("\t", "    ")
    return "".join(_PRIVATE.get(c, c) if "" <= c <= "" else c for c in text)


def _marker_font(ch: str) -> str:
    return "Arial" if ch in "•·●○◦▪■□–—-*" else ("Segoe UI Symbol" if ch in _BULLETS else "Arial")


# ----------------------------------------------------------------------------
# 图片
# ----------------------------------------------------------------------------


def fix_jpeg(data: bytes) -> bytes:
    """JFIF 里的分辨率为 0 时,python-docx 会除零报错:改成 96 dpi。"""
    if data[:2] == b"\xff\xd8" and data[2:4] == b"\xff\xe0" and data[6:11] == b"JFIF\x00":
        b = bytearray(data)
        b[13] = 1
        b[14:16] = (96).to_bytes(2, "big")
        b[16:18] = (96).to_bytes(2, "big")
        return bytes(b)
    return data


def image_bytes(doc, xref: int):
    """提取 PDF 里的一张图片(带透明通道的合成 PNG;普通 JPEG 保持原样)。"""
    try:
        smask = doc.xref_get_key(xref, "SMask")
        info = doc.extract_image(xref)
        if info and info.get("ext") in ("jpg", "jpeg") and smask[0] == "null":
            return fix_jpeg(info["image"])
        pix = fitz.Pixmap(doc, xref)
        if pix.n - pix.alpha > 3:
            pix = fitz.Pixmap(fitz.csRGB, pix)
        if smask[0] == "xref" and not pix.alpha:
            try:
                pix = fitz.Pixmap(pix, fitz.Pixmap(doc, int(smask[1].split()[0])))
            except Exception:  # noqa: BLE001
                pass
        return pix.tobytes("png")
    except Exception:  # noqa: BLE001
        return None


def _png_or_jpeg(data: bytes) -> bytes:
    """python-docx 读不了的图片,用 PIL 重新编码。"""
    img = Image.open(io.BytesIO(data))
    img.load()
    buf = io.BytesIO()
    if img.mode in ("RGBA", "LA", "P"):
        img.convert("RGBA").save(buf, "PNG")
    else:
        img.convert("RGB").save(buf, "JPEG", quality=92, dpi=(96, 96))
    return buf.getvalue()


def _crop_png_or_jpeg(page_img: Image.Image, rect) -> bytes:
    """从 150dpi 的整页渲染图里截取 rect(pt)。小图用 PNG,大图用 JPEG。"""
    box = tuple(max(0, int(round(v * 150 / 72))) for v in rect)
    if box[2] - box[0] < 2 or box[3] - box[1] < 2:  # 图片在页面外 / 太小,没东西可截
        return None
    img = page_img.crop(box).convert("RGB")
    buf = io.BytesIO()
    if rect.width * rect.height < 200 * 200:
        img.save(buf, "PNG")
    else:
        img.save(buf, "JPEG", quality=90, dpi=(150, 150))
    return buf.getvalue()


def _matches_render(data: bytes, rect, page_img: Image.Image) -> bool:
    """提取出来的原图,是否和页面上这个位置实际渲染出来的样子一致(不一致 = 被裁剪或被别的东西盖住)。"""
    try:
        raw = Image.open(io.BytesIO(data))
        raw.load()
        rgba = raw.convert("RGBA")
        flat = Image.new("RGB", rgba.size, (255, 255, 255))
        flat.paste(rgba, mask=rgba.split()[-1])
        a = flat.resize((48, 48))
        box = tuple(max(0, int(round(v * 150 / 72))) for v in rect)
        b = page_img.crop(box).convert("RGB").resize((48, 48))
        diff = sum(abs(x - y) for p, q in zip(a.getdata(), b.getdata()) for x, y in zip(p, q)) / (48 * 48 * 3)
        return diff < 22
    except Exception:  # noqa: BLE001
        return True


def _is_blank(img: Image.Image) -> bool:
    """整张图几乎全白(细线也算内容,所以不能先缩小再判断)。"""
    return img.convert("L").point(lambda v: 255 if v < 240 else 0).getbbox() is None


def _upright(transform) -> bool:
    a, b, c, d = transform[:4]
    return a > 0 and d > 0 and abs(b) < 1e-3 * abs(a) + 1e-6 and abs(c) < 1e-3 * abs(d) + 1e-6


def _overlap(a, b) -> float:
    """a 被 b 盖住的比例。"""
    inter = a & b
    if inter.is_empty or a.get_area() <= 0:
        return 0.0
    return inter.get_area() / a.get_area()


# ----------------------------------------------------------------------------
# 版面分析
# ----------------------------------------------------------------------------


def _rect(r, rm):
    return (fitz.Rect(r) * rm).normalize()


def _block_rows(block, rm):
    """一个文本块 → 若干"行"(同一基线上的几段文字合成一行,例如 "•" 和后面的文字)。"""
    lines = []
    for ln in block["lines"]:
        spans = [s for s in ln["spans"] if s["text"]]
        if not spans or not "".join(s["text"] for s in spans).strip():
            continue
        for s in spans:
            s["_r"] = _rect(s["bbox"], rm)
        lines.append(dict(bbox=_rect(ln["bbox"], rm), spans=spans))
    rows = []
    for ln in lines:
        for r in rows[-3:]:
            ov = min(r["bbox"].y1, ln["bbox"].y1) - max(r["bbox"].y0, ln["bbox"].y0)
            if ov > 0.55 * min(r["bbox"].height, ln["bbox"].height) and (
                    ln["bbox"].x0 >= r["bbox"].x1 - 2 or ln["bbox"].x1 <= r["bbox"].x0 + 2):
                r["spans"] = sorted(r["spans"] + ln["spans"], key=lambda s: s["_r"].x0)
                r["bbox"] = r["bbox"] | ln["bbox"]
                break
        else:
            rows.append(ln)
    for r in rows:
        r["size"] = max(s["size"] for s in r["spans"])
        r["base"] = sorted(fitz.Point(s["origin"]).transform(rm).y for s in r["spans"])[len(r["spans"]) // 2]
        r["marker"] = _detect_marker(r)
    return rows


def _detect_marker(row):
    """行首是项目符号 / 编号、后面隔着一段空隙才是正文 → 列表项。返回 dict 并把行里的符号 span 拿掉。"""
    sp = row["spans"]
    i = 0
    while i < len(sp) and not clean_text(sp[i]["text"]).strip():
        i += 1
    if i >= len(sp) - 1:
        return None
    first = sp[i]
    t = clean_text(first["text"]).strip()
    if not t or len(t) > 5:
        return None
    is_bullet = all(c in _BULLETS for c in t)
    if not (is_bullet or _NUM_MARK.match(t)):
        return None
    rest = [s for s in sp[i + 1:] if clean_text(s["text"]).strip()]
    if not rest:
        return None
    nxt = rest[0]
    gap = nxt["_r"].x0 - first["_r"].x1
    spaced = clean_text(nxt["text"]).startswith(" ") or any(
        not clean_text(s["text"]).strip() for s in sp[i + 1:sp.index(nxt)])
    if gap < 0.2 * row["size"] and not spaced:
        return None
    text_x0 = nxt["_r"].x0 + (0.28 * nxt["size"] if clean_text(nxt["text"]).startswith(" ") and gap < 0.2 * row["size"] else 0)
    row["spans"] = sp[sp.index(nxt):]
    first_txt = clean_text(row["spans"][0]["text"]).lstrip()
    row["spans"][0] = dict(row["spans"][0], text=first_txt)
    return dict(char=t, x0=first["_r"].x0, text_x0=text_x0, size=first["size"], bullet=is_bullet)


def _vector_dots(page, rm):
    dots = []
    try:
        for d in page.get_drawings():
            r = d["rect"]
            if 2 <= r.width <= 12 and 2 <= r.height <= 12 and abs(r.width - r.height) <= 3 and d.get("fill") is not None:
                circle = bool(d["items"]) and d["items"][0][0] == "c"
                dots.append((_rect(r, rm), "•" if circle else "▪"))
    except Exception:  # noqa: BLE001
        pass
    return dots


def _para_from_rows(rows):
    bb = rows[0]["bbox"]
    for r in rows[1:]:
        bb = bb | r["bbox"]
    return dict(kind="para", bbox=bb, rows=rows)


def _text_left(para) -> float:
    r0 = para["rows"][0]
    return r0["marker"]["text_x0"] if r0["marker"] else r0["bbox"].x0


def _eff_width(row) -> float:
    return row["bbox"].x1 - (row["marker"]["text_x0"] if row["marker"] else row["bbox"].x0)


def merge_paragraphs(els: list) -> list:
    """有的 PDF 把每一行都当成独立的文本块。把相邻的、左对齐相同、字号相同、行距紧凑、
    上一行排满了的块合并回同一个自然段(列表项的续行并入列表项)。"""
    out = []
    for e in els:
        if out and e["kind"] == "para" and out[-1]["kind"] == "para" and not e["rows"][0]["marker"]:
            a = out[-1]
            la, lb = a["rows"][-1], e["rows"][0]
            h = la["bbox"].height
            gap = lb["bbox"].y0 - la["bbox"].y1
            a_left = la["bbox"].x0 if len(a["rows"]) > 1 or not a["rows"][0]["marker"] else a["rows"][0]["marker"]["text_x0"]
            wmax = max(_eff_width(r) for r in a["rows"] + e["rows"])
            # 左边对齐,或者 a 只有一行、比下面缩进了几个字(首行缩进的段落常被拆成"首行"和"其余各行"两块)
            first_indent = (len(a["rows"]) == 1 and not a["rows"][0]["marker"]
                            and 0 < la["bbox"].x0 - lb["bbox"].x0 <= 4.5 * lb["size"])
            if (-0.5 * h <= gap <= 0.6 * h and lb["bbox"].y0 > la["bbox"].y0
                    and (abs(a_left - lb["bbox"].x0) < 2.5 or first_indent)
                    and abs(la["size"] - lb["size"]) < 1 and _eff_width(la) >= 0.82 * wmax):
                a["rows"] = a["rows"] + e["rows"]
                a["bbox"] = a["bbox"] | e["bbox"]
                continue
        out.append(e)
    return out


def _hidden_ratio(page) -> float:
    try:
        tot = hid = 0
        for t in page.get_texttrace():
            n = len(t.get("chars", ()))
            tot += n
            if t.get("type") == 3:
                hid += n
        return hid / tot if tot else 0.0
    except Exception:  # noqa: BLE001
        return 0.0


def ocr_elements(lines: list) -> list:
    """识别出的行 [{'text','bbox'(pt)}] 合并成段落元素。"""
    lines = sorted(lines, key=lambda l: (l["bbox"][1], l["bbox"][0]))
    font = E.font_for_text("yahei", False, False, "中")
    groups = []
    for ln in lines:
        x0, y0, x1, y1 = ln["bbox"]
        h = y1 - y0
        if groups:
            last = groups[-1][-1]["bbox"]
            if 0 <= y0 - last[3] < 0.9 * h and abs(x0 - last[0]) < 1.5 * h:
                groups[-1].append(ln)
                continue
        groups.append([ln])
    els = []
    for g in groups:
        rows = []
        for ln in g:
            geo = E.ocr_line_geometry(ln["text"], ln["bbox"], font)
            size = geo[0] if geo else 11.0
            r = fitz.Rect(*ln["bbox"])
            span = dict(text=ln["text"], size=size, flags=0, color=0, font="MicrosoftYaHei", origin=(r.x0, r.y1), _r=r)
            rows.append(dict(bbox=r, spans=[span], size=size, base=r.y1, marker=None))
        els.append(_para_from_rows(rows))
    return els


def _hex(color) -> str | None:
    if not color:
        return None
    return "#%02x%02x%02x" % tuple(int(round(max(0.0, min(1.0, v)) * 255)) for v in color[:3])


def _fill_rects(page, rm) -> list:
    """页面上有底色的矩形区域 [(rect, '#rrggbb')](不含白色)。"""
    out = []
    try:
        for d in page.get_drawings():
            col = _hex(d.get("fill"))
            r = _rect(d["rect"], rm)
            if col and col != "#ffffff" and r.width > 3 and r.height > 3:
                out.append((r, col))
    except Exception:  # noqa: BLE001
        pass
    return out


def _cell_fill(cell, fills) -> str | None:
    best = None
    for r, col in fills:
        if _overlap(cell, r) > 0.8 and r.get_area() <= 6 * cell.get_area():
            if best is None or r.get_area() < best[0]:
                best = (r.get_area(), col)
    return best[1] if best else None


def _mark_underlines(paras, page, rm) -> list:
    """细横线紧贴在文字基线下方 → 这段文字加下划线;穿过文字中部 → 删除线。返回用掉的细线(要从背景层里去掉)。"""
    rules = []
    try:
        for d in page.get_drawings():
            r = _rect(d["rect"], rm)
            if r.height <= 1.8 and r.width >= 4:
                rules.append(r)
    except Exception:  # noqa: BLE001
        return []
    used = set()
    for p in paras:
        for row in p["rows"]:
            for s in row["spans"]:
                sr = s.get("_r")
                if sr is None or not s["text"].strip():
                    continue
                base = fitz.Point(s["origin"]).transform(rm).y
                size = s["size"]
                for k, r in enumerate(rules):
                    ov = min(r.x1, sr.x1) - max(r.x0, sr.x0)
                    if ov < 0.6 * sr.width:
                        continue
                    cy = (r.y0 + r.y1) / 2
                    if base - 0.1 * size <= cy <= base + 0.35 * size:
                        s["_u"] = True
                        used.add(k)
                    elif base - 0.6 * size <= cy <= base - 0.15 * size:
                        s["_x"] = True
                        used.add(k)
    return [rules[k] for k in used]


def _graphics_layer(page, erase) -> bytes | None:
    """只渲染页面上的矢量图形(去掉文字和图片),并把 erase 里的区域涂白。整页空白返回 None。"""
    try:
        if not page.get_drawings():
            return None
        tmp = fitz.open()
        tmp.insert_pdf(page.parent, from_page=page.number, to_page=page.number)
        tp = tmp[0]
        tp.add_redact_annot((fitz.Rect(tp.rect) * tp.derotation_matrix).normalize(), fill=False, cross_out=False)
        try:
            tp.apply_redactions(images=fitz.PDF_REDACT_IMAGE_REMOVE, graphics=fitz.PDF_REDACT_LINE_ART_NONE)
        except (TypeError, AttributeError):
            tp.apply_redactions(images=fitz.PDF_REDACT_IMAGE_REMOVE)
        pm = tp.get_pixmap(dpi=150, alpha=False)
        img = Image.frombytes("RGB", (pm.width, pm.height), pm.samples)
        tmp.close()
        from PIL import ImageDraw
        draw = ImageDraw.Draw(img)
        k = 150 / 72
        for r in erase:
            draw.rectangle([r.x0 * k, r.y0 * k, r.x1 * k, r.y1 * k], fill=(255, 255, 255))
        if _is_blank(img):
            return None
        buf = io.BytesIO()
        img.save(buf, "PNG", optimize=True)
        return buf.getvalue()
    except Exception:  # noqa: BLE001
        return None


def _merge_same_row(els: list) -> list:
    """同一行上左右分开的几段文字(例如"姓名:……  日期:……")合成一个段落,中间用制表符隔开,
    否则 Word 只能把它们上下排,后面的内容都会被挤下去。"""
    out = []
    for e in els:
        if (out and e["kind"] == "para" and out[-1]["kind"] == "para" and len(e["rows"]) == 1
                and len(out[-1]["rows"]) == 1 and not e["rows"][0]["marker"]):
            ra, rb = out[-1]["rows"][0], e["rows"][0]
            if abs(ra["base"] - rb["base"]) < 2.5 and rb["bbox"].x0 > ra["bbox"].x1 + 4:
                tab = dict(text="\t", _tab=True, size=rb["size"], font=rb["spans"][0]["font"], flags=0, color=0,
                           origin=rb["spans"][0]["origin"], _r=fitz.Rect(ra["bbox"].x1, rb["bbox"].y0, rb["bbox"].x0, rb["bbox"].y1))
                ra["spans"] = ra["spans"] + [tab] + rb["spans"]
                ra["tabs"] = ra.get("tabs", []) + [rb["bbox"].x0]
                ra["bbox"] = ra["bbox"] | rb["bbox"]
                out[-1]["bbox"] = out[-1]["bbox"] | e["bbox"]
                continue
        out.append(e)
    return out


def analyze_page(page, fonts: Fonts, ocr_lines=None, textless=None):
    """把一页分析成元素列表。textless() 返回"去掉文字后的整页渲染图"(PIL,150dpi),需要时才会调用。"""
    rm = page.rotation_matrix
    W, H = page.rect.width, page.rect.height
    PAGE = fitz.Rect(0, 0, W, H)
    hidden = _hidden_ratio(page) >= 0.8
    has_text = E.page_has_text(page)
    scan_lines = None
    if hidden:  # 带不可见文字层的扫描件:用文字层的行重新排段落,页面大图丢掉
        scan_lines = [{"text": l["text"], "bbox": list(l["bbox"])} for l in E.page_lines(page)]
    elif not has_text and ocr_lines:
        scan_lines = ocr_lines
    els, table_rects, notes = [], [], []

    # ---- 表格
    if has_text and not hidden:
        try:
            # 默认策略会把 Word 表格里单元格底色的边缘当成多余的列线;严格按"线"检测更准,检测不到再退回默认策略
            found = page.find_tables(vertical_strategy="lines_strict", horizontal_strategy="lines_strict").tables
            all_spans = [s for b in page.get_text("dict", flags=E._TEXT_FLAGS)["blocks"] if b["type"] == 0
                         for ln in b["lines"] for s in ln["spans"] if s["text"].strip()]
            for s in all_spans:
                s["_r"] = _rect(s["bbox"], rm)
            fills = _fill_rects(page, rm)
            for t in (found or page.find_tables().tables):
                r = _rect(t.bbox, rm)
                rows = t.extract()
                if rows and _real_table(t, rows, all_spans):
                    cells = [[_rect(c, rm) if c else None for c in row.cells] for row in t.rows]
                    shade = [[_cell_fill(c, fills) if c else None for c in row] for row in cells]
                    spans = [s for s in all_spans if r.contains(fitz.Point((s["_r"].x0 + s["_r"].x1) / 2,
                                                                           (s["_r"].y0 + s["_r"].y1) / 2))]
                    els.append(dict(kind="table", bbox=r, rows=rows, cols=_table_cols(t, rm), cells=cells,
                                    shade=shade, spans=spans))
                    table_rects.append(r)
        except Exception:  # noqa: BLE001
            pass

    # ---- 文字
    if scan_lines:
        els += ocr_elements(scan_lines)
    elif has_text:
        dots = _vector_dots(page, rm)
        for b in page.get_text("dict", flags=E._TEXT_FLAGS, sort=True)["blocks"]:
            if b["type"] != 0:
                continue
            bb = _rect(b["bbox"], rm)
            if any(_overlap(bb, t) > 0.6 for t in table_rects):
                continue
            rows = _block_rows(b, rm)
            if not rows:
                continue
            for row in rows:  # 矢量圆点项目符号
                if row["marker"] is None:
                    for dr, ch in dots:
                        cy = (dr.y0 + dr.y1) / 2
                        if row["bbox"].y0 <= cy <= row["bbox"].y1 and dr.x1 <= row["bbox"].x0 + 1 and row["bbox"].x0 - dr.x1 <= 40:
                            row["marker"] = dict(char=ch, x0=dr.x0, text_x0=row["bbox"].x0, size=row["size"], bullet=True)
                            break
            groups, cur = [], []  # 一个块里有多个列表项时,每个列表项是一个段落
            for row in rows:
                if row["marker"] and cur:
                    groups.append(cur)
                    cur = []
                cur.append(row)
            groups.append(cur)
            els += [_para_from_rows(g) for g in groups]

    # ---- 下划线 / 删除线:PDF 里是画在文字下面 / 中间的细线,改成真正的文字格式
    used_rules = _mark_underlines([e for e in els if e["kind"] == "para"], page, rm) if has_text and not hidden else []

    # ---- 矢量图(图表、卡片):里面文字不多时,连文字一起渲染成图片,外观最准
    floats, pictures = [], []
    if has_text and not hidden and page.rotation == 0:
        try:
            clusters = page.cluster_drawings()
        except Exception:  # noqa: BLE001
            clusters = []
        for c in clusters:
            r = fitz.Rect(c)
            if r.width < 24 or r.height < 24 or r.get_area() > 0.5 * W * H:
                continue
            if any(_overlap(r, t) > 0.3 for t in table_rects):
                continue
            if len(page.get_text(clip=r).strip()) > 120:
                continue  # 文字多的色块:交给下面的背景层,文字保持可编辑
            png = page.get_pixmap(dpi=150, clip=r, alpha=False).tobytes("png")
            if _is_blank(Image.open(io.BytesIO(png))):
                continue
            els = [e for e in els if not (e["kind"] == "para" and _overlap(e["bbox"], r + (-2, -2, 2, 2)) > 0.8)]
            els = [e for e in els if not (e["kind"] == "image" and _overlap(e["bbox"], r) > 0.7)]
            els.append(dict(kind="image", bbox=r, data=png, float_=r.get_area() > 0.3 * W * H, text_in=True))
            pictures.append(r)

    # ---- 背景层:只保留矢量图形(色块、竖条、装饰线、边框)渲染成一张整页图,衬在最底下
    if has_text and not hidden:
        bg = _graphics_layer(page, [r + (-2, -2, 2, 2) for r in table_rects] + pictures
                             + [r + (-1, -1, 1, 1) for r in used_rules])
        if bg is not None:
            floats.append(dict(kind="image", bbox=fitz.Rect(PAGE), data=bg, float_=True, behind=True, background=True))

    # ---- 图片
    try:
        infos = page.get_image_info(xrefs=True)
    except Exception:  # noqa: BLE001
        infos = []
    for info in infos:
        r = _rect(info["bbox"], rm)
        if r.width < 4 or r.height < 4:
            continue
        if any(e["kind"] == "image" and e.get("text_in") and _overlap(r, e["bbox"]) > 0.7 for e in els):
            continue  # 已经包含在某块矢量图里了
        if (hidden or scan_lines) and r.get_area() > 0.6 * W * H:
            continue  # 扫描底图:用识别的文字代替
        data = None
        inside = r.x0 >= -1 and r.y0 >= -1 and r.x1 <= W + 1 and r.y1 <= H + 1
        if info.get("xref") and inside and _upright(info.get("transform", (1, 0, 0, 1))):
            data = image_bytes(page.parent, info["xref"])
            # 原图和页面上实际看到的不一致(被裁剪、被遮住)时,改用页面渲染的截取
            if data is not None and textless is not None and not _matches_render(data, r, textless()):
                data = None
        if data is None and textless is not None:  # 没有编号 / 倒置翻转 / 被裁剪 / 提取失败:直接从页面渲染截取
            data = _crop_png_or_jpeg(textless(), r & PAGE)
        if data is None:
            if (r & PAGE).get_area() > 25:  # 页面外的碎图不算丢失
                notes.append("有图片无法提取")
            continue
        els.append(dict(kind="image", bbox=r & PAGE if not inside else r, data=data))

    # ---- 图片是内联还是浮动:和文字上下有重叠(图文并排 / 文字压在图上)、背景大图 → 固定在页面位置、衬于文字下方
    others = [e for e in els if e["kind"] in ("para", "table")]
    for e in els:
        if e["kind"] != "image":
            continue
        bb = e["bbox"]
        shares_rows = any(min(bb.y1, o["bbox"].y1) - max(bb.y0, o["bbox"].y0) > 0.2 * min(bb.height, o["bbox"].height)
                          for o in others)
        bleeds = bb.x0 < 12 or bb.y0 < 12 or bb.x1 > W - 12 or bb.y1 > H - 12  # 贴着页边的图(通栏头图等)放不进页边距
        e["float_"] = bool(e.get("float_") or shares_rows or bleeds or bb.get_area() > 0.6 * W * H)
        e.setdefault("behind", True)
    flow = [e for e in els if not (e["kind"] == "image" and e["float_"])]
    floating = [e for e in els if e["kind"] == "image" and e["float_"]] + floats
    flow.sort(key=lambda e: (round(e["bbox"].y0 / 4), e["bbox"].x0))
    floating.sort(key=lambda e: -e["bbox"].get_area())  # 大的在下面
    return merge_paragraphs(_merge_same_row(_merge_markers(flow))), floating, notes


def _real_table(t, rows, spans) -> bool:
    """排除"假表格":封面上的装饰线、色块边框也会被检测成表格。
    真表格至少有两行、每行两个以上有内容的单元格,而且列线不会从一段文字的中间穿过去。"""
    filled_rows = sum(1 for row in rows if sum(1 for c in row if c and str(c).strip()) >= 2)
    if filled_rows < 2:
        return False
    try:
        x0, y0, x1, y1 = t.bbox
        edges = sorted({round(c[0], 1) for row in t.rows for c in row.cells if c} - {round(x0, 1)})
        for s in spans:
            sx0, sy0, sx1, sy1 = s["bbox"]
            if sy1 < y0 or sy0 > y1 or sx1 < x0 or sx0 > x1:
                continue
            cut = [e for e in edges if sx0 + 2 < e < sx1 - 2]
            if cut and len(s["text"].strip()) > 1:
                return False
    except Exception:  # noqa: BLE001
        pass
    return True


def _merge_markers(els: list) -> list:
    """项目符号和正文分别在两个文本块里(同一行、符号在左)时,把符号并进正文,变成列表项。"""
    out = []
    for e in els:
        if out and e["kind"] == "para" and out[-1]["kind"] == "para":
            a = out[-1]
            if len(a["rows"]) == 1 and not a["rows"][0]["marker"]:
                ar, br = a["rows"][0], e["rows"][0]
                t = clean_text("".join(s["text"] for s in ar["spans"])).strip()
                is_bullet = bool(t) and all(c in _BULLETS for c in t)
                if ((is_bullet or (_NUM_MARK.match(t) if t else None)) and len(t) <= 5 and not br["marker"]
                        and min(ar["bbox"].y1, br["bbox"].y1) - max(ar["bbox"].y0, br["bbox"].y0) > 0.5 * min(ar["bbox"].height, br["bbox"].height)
                        and 0 <= br["bbox"].x0 - ar["bbox"].x1 < 60):
                    br["marker"] = dict(char=t, x0=ar["bbox"].x0, text_x0=br["bbox"].x0, size=ar["size"], bullet=is_bullet)
                    e["bbox"] = e["bbox"] | a["bbox"]
                    out[-1] = e
                    continue
        out.append(e)
    return out


def _table_cols(t, rm):
    """表格各列的左右边界(用第一行单元格的位置)和每行的高度。"""
    out = {"xs": None, "heights": None}
    try:
        cells = [c for c in t.rows[0].cells if c]
        xs = sorted({round(c[0], 1) for c in cells} | {round(c[2], 1) for c in cells})
        out["xs"] = xs if len(xs) >= 2 else None
        out["heights"] = [r.bbox[3] - r.bbox[1] for r in t.rows]
    except Exception:  # noqa: BLE001
        pass
    return out


# ----------------------------------------------------------------------------
# 写出 Word
# ----------------------------------------------------------------------------


def _set_font(run, family: str):
    from docx.oxml.ns import qn
    run.font.name = family
    rf = run._element.rPr.rFonts
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(attr), family)


def _half(size: float) -> float:
    """Word 的字号只能是半磅的整数倍。就近取,但不能比原来大(大一点点,排满的一行最后一个字就会被挤到下一行)。"""
    r = round(size * 2) / 2
    return r - 0.5 if r - size > 0.05 and r > 1 else r


def _add_runs(p, spans, page, fonts: Fonts, scale: float):
    from docx.shared import Pt, RGBColor
    groups = []  # [(样式键, 文字)];制表符单独一组,样式键为 None
    for s in spans:
        if s.get("_tab"):
            groups.append((None, "\t"))
            continue
        fam, bold, italic = fonts.info(page, s)
        key = (fam, _half(s["size"] * scale), bold, italic, s["color"], bool(s["flags"] & 1),
               bool(s.get("_u")), bool(s.get("_x")))
        text = clean_text(s["text"])
        if groups and groups[-1][0] == key:
            groups[-1] = (key, groups[-1][1] + text)
        else:
            groups.append((key, text))
    for key, text in groups:
        if key is None:
            p.add_run("\t")
            continue
        fam, size, bold, italic, color, sup, under, strike = key
        if not text:
            continue
        run = p.add_run(text)
        _set_font(run, fam)
        run.font.size = Pt(max(size, 1))  # size 已经乘过缩放并取到半磅
        run.font.bold = bold or None
        run.font.italic = italic or None
        if under:
            run.font.underline = True
        if strike:
            run.font.strike = True
        if sup:
            run.font.superscript = True
        if color:
            run.font.color.rgb = RGBColor((color >> 16) & 255, (color >> 8) & 255, color & 255)


def _is_cjk(ch: str) -> bool:
    o = ord(ch)
    return (0x2E80 <= o <= 0xA4CF) or (0xF900 <= o <= 0xFAFF) or (0xFE30 <= o <= 0xFE4F) or (0xFF00 <= o <= 0xFFEF)


def _join(prev: str, nxt: str) -> str:
    return "" if prev and nxt and (_is_cjk(prev[-1]) or _is_cjk(nxt[0])) else " "


def para_metrics(e):
    """段落在 Word 里的固定行高、"期望的顶端位置"和总高度(都按 PDF 坐标,pt)。

    Word 的固定行高会把多出来的空白放在文字上方(文字贴着行底),所以段落顶端要比 PDF 里字体框的顶端
    高出 (行高 - 字体自然高度) 左右,这样文字基线才落在 PDF 里的原位置。"""
    rows = e["rows"]
    biggest = max(r["size"] for r in rows)
    pitch = (rows[-1]["base"] - rows[0]["base"]) / (len(rows) - 1) if len(rows) > 1 else rows[0]["bbox"].height
    exact = max(pitch, 1.08 * biggest) if pitch > 4 else None
    line = exact or biggest * 1.2
    return exact, rows[0]["base"] - line + 0.24 * biggest, line * len(rows)


def write_paragraph(d, page, fonts, e, left_m, text_w, W, gap, scale):
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.shared import Pt

    rows, bb = e["rows"], e["bbox"]
    first = rows[0]
    marker = first["marker"]
    text_left = _text_left(e)
    cont_left = min([r["bbox"].x0 for r in rows[1:]] + [text_left]) if len(rows) > 1 else text_left
    widths = [_eff_width(r) for r in rows]
    block_w = max(widths)
    reflow = len(rows) >= 2 and block_w >= 0.35 * text_w and all(w >= 0.82 * block_w for w in widths[:-1])

    p = d.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(gap * scale)
    pf.space_after = Pt(0)
    mids = [(r["bbox"].x0 + r["bbox"].x1) / 2 for r in rows]
    area_mid = left_m + text_w / 2
    centered = (not marker and bb.x0 - left_m > 8 and
                all(abs(m - W / 2) < 3 or abs(m - area_mid) < 3 for m in mids) and
                (len(rows) == 1 or len({round(r["bbox"].x0) for r in rows}) > 1))
    tabs = first.get("tabs")
    if centered and not tabs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif (not marker and not tabs and len(rows) == 1 and (left_m + text_w) - bb.x1 < 9 and bb.x0 - left_m > 0.4 * text_w):
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        if marker:
            pf.left_indent = Pt(max(text_left - left_m, 0) * scale)
            pf.first_line_indent = Pt(-max(text_left - marker["x0"], 0) * scale)
            pf.tab_stops.add_tab_stop(Pt(max(text_left - left_m, 0) * scale))
        else:
            if cont_left - left_m > 2.5:
                pf.left_indent = Pt((cont_left - left_m) * scale)
            indent1 = first["bbox"].x0 - cont_left
            if len(rows) > 1 and abs(indent1) > 2.5:
                pf.first_line_indent = Pt(indent1 * scale)
        for x in tabs or ():  # 同一行左右分开的几段文字:按原位置设制表位
            pf.tab_stops.add_tab_stop(Pt(max(x - left_m, 0) * scale))
        # 两端对齐:除最后一行外,每行的右边缘都对得很齐
        rights = [r["bbox"].x1 for r in rows[:-1]]
        if reflow and len(rights) >= 1 and max(rights) - min(rights) < 1.5 and bb.width > 0.5 * text_w:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            # 右边留一点余量:Word 排中文标点比 PDF 稍宽,中文段落多留一个字宽,否则行尾的字会被挤到下一行
            body = "".join(s["text"] for r in rows for s in r["spans"])
            cjk = sum(1 for c in body if _is_cjk(c)) / max(len(body), 1)
            slack = max(r["size"] for r in rows) if cjk > 0.3 else 1.5
            pf.right_indent = Pt(max(0.0, (left_m + text_w) - max(rights) - slack) * scale)
    # 固定行高,和 PDF 一致。Word 默认的行高比 PDF 大,不固定的话内容会越排越长、溢出到下一页
    exact, _top, _h = para_metrics(e)
    if exact:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(exact * scale)  # 不低于字号的 1.08 倍,避免字被裁掉
    if marker:  # 项目符号 / 编号 + 制表符
        ch = marker["char"]
        run = p.add_run(ch)
        _set_font(run, _marker_font(ch) if marker["bullet"] else fonts.family(page, first["spans"][0]["font"]))
        run.font.size = Pt(max(marker["size"] * scale, 1))
        p.add_run("\t")
    for li, row in enumerate(rows):
        if li:
            if reflow:
                prev_t = clean_text(rows[li - 1]["spans"][-1]["text"])
                nxt_t = clean_text(row["spans"][0]["text"])
                p.add_run(_join(prev_t, nxt_t))
            else:
                p.add_run().add_break()
        _add_runs(p, row["spans"], page, fonts, scale)


def write_table(d, page, fonts, e, text_w, scale):
    from docx.shared import Pt

    rows = e["rows"]
    ncols = max(len(r) for r in rows)
    # 表格里的字体和字号:取表格区域里最常见的
    fam, size, bold_head = "Microsoft YaHei", 10.0, False
    try:
        spans = [s for b in page.get_text("dict", clip=fitz.Rect(e["bbox"]) * page.derotation_matrix, flags=E._TEXT_FLAGS)["blocks"]
                 if b["type"] == 0 for ln in b["lines"] for s in ln["spans"] if s["text"].strip()]
        if spans:
            c = Counter((fonts.family(page, s["font"]), round(s["size"] * 2) / 2) for s in spans)
            (fam, size), _n = c.most_common(1)[0]
            first_y = min(s["bbox"][1] for s in spans)
            heads = [s for s in spans if s["bbox"][1] < first_y + size * 1.5]
            bold_head = bool(heads) and sum(1 for s in heads if E.guess_font_style(s["font"], s["flags"])[1]) > len(heads) / 2
    except Exception:  # noqa: BLE001
        pass
    from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_ALIGN_VERTICAL

    t = d.add_table(rows=len(rows), cols=ncols)
    t.style = "Table Grid"
    t.autofit = False
    cols = (e.get("cols") or {}).get("xs")
    heights = (e.get("cols") or {}).get("heights")
    widths = None
    if cols and len(cols) - 1 == ncols:
        widths = [(cols[i + 1] - cols[i]) for i in range(ncols)]
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    cells = e.get("cells") or []
    shade = e.get("shade") or []
    tspans = e.get("spans") or []
    for ri, row in enumerate(rows):
        single = all("\n" not in str(c or "").strip() for c in row)
        if heights and ri < len(heights):  # 行高和原表格一致;每个格子都只有一行字时用固定行高,不让 Word 撑高
            t.rows[ri].height = Pt(heights[ri] * scale)
            t.rows[ri].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY if single else WD_ROW_HEIGHT_RULE.AT_LEAST
        for ci in range(ncols):
            cell = t.cell(ri, ci)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = Pt(min(widths[ci], text_w) * scale)
            crect = cells[ri][ci] if ri < len(cells) and ci < len(cells[ri]) else None
            fill = shade[ri][ci] if ri < len(shade) and ci < len(shade[ri]) else None
            if fill:  # 单元格底色
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), fill[1:].upper())
                cell._tc.get_or_add_tcPr().append(shd)
            inner = []
            if crect is not None:
                inner = [s for s in tspans if crect.contains(fitz.Point((s["_r"].x0 + s["_r"].x1) / 2, (s["_r"].y0 + s["_r"].y1) / 2))]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if inner:  # 用单元格里原来的文字片段:保留粗体、斜体、颜色、字号
                inner.sort(key=lambda s: (round(fitz.Point(s["origin"]).transform(page.rotation_matrix).y), s["_r"].x0))
                prev = None
                for s in inner:
                    if prev is not None:
                        same_line = abs(s["_r"].y0 - prev["_r"].y0) < 0.5 * s["size"]
                        if not same_line:
                            p.add_run(_join(clean_text(prev["text"]), clean_text(s["text"])))
                        elif s["_r"].x0 - prev["_r"].x1 > 0.25 * s["size"]:
                            p.add_run(" ")
                    _add_runs(p, [s], page, fonts, scale)
                    prev = s
                tx0 = min(s["_r"].x0 for s in inner)
                tx1 = max(s["_r"].x1 for s in inner)
                if abs((tx0 + tx1) / 2 - (crect.x0 + crect.x1) / 2) < 2.5 and tx0 - crect.x0 > 4:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif crect.x1 - tx1 < 5 and tx0 - crect.x0 > 10:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                p.add_run(clean_text(((row[ci] if ci < len(row) else "") or "").strip()))
                for r in p.runs:
                    r.font.size = Pt(size * scale)
                    _set_font(r, fam)
                    if bold_head and ri == 0:
                        r.font.bold = True


def _add_picture(paragraph, data: bytes, w: float, h: float, warnings: list):
    from docx.shared import Pt
    run = paragraph.add_run()
    for attempt in (data, None):
        try:
            payload = attempt if attempt is not None else _png_or_jpeg(data)
            return run.add_picture(io.BytesIO(payload), width=Pt(w), height=Pt(h))
        except Exception:  # noqa: BLE001
            continue
    warnings.append("有图片无法写入 Word")
    return None


def write_inline_image(d, e, left_m, text_w, W, gap, scale, warnings):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    bb = e["bbox"]
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(gap * scale)
    p.paragraph_format.space_after = Pt(0)
    if abs((bb.x0 + bb.x1) / 2 - W / 2) < 3 and bb.x0 - left_m > 8:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif bb.x0 - left_m > 3:
        p.paragraph_format.left_indent = Pt(min((bb.x0 - left_m) * scale, text_w * scale * 0.6))
    w = min(bb.width, text_w)
    _add_picture(p, e["data"], w * scale, bb.height * (w / bb.width) * scale, warnings)


def write_floating(d, floating, scale, warnings):
    """页面上所有浮动图片放在这一页最前面的一个小段落里,按位置固定在页面上并衬于文字下方。"""
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    from docx.shared import Pt

    if not floating:
        return
    p = d.add_paragraph()
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(1)
    for z, e in enumerate(floating):
        bb = e["bbox"]
        shape = _add_picture(p, e["data"], bb.width * scale, bb.height * scale, warnings)
        if shape is None:
            continue
        inline = shape._inline
        emu = lambda v: int(v * scale * 12700)  # noqa: E731
        anchor = parse_xml(
            f'<wp:anchor {nsdecls("wp", "a", "pic")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
            f'relativeHeight="{251658240 + z}" behindDoc="{1 if e.get("behind", True) else 0}" locked="0" '
            f'layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
            f'<wp:positionH relativeFrom="page"><wp:posOffset>{emu(bb.x0)}</wp:posOffset></wp:positionH>'
            f'<wp:positionV relativeFrom="page"><wp:posOffset>{emu(bb.y0)}</wp:posOffset></wp:positionV>'
            f'<wp:extent cx="{emu(bb.width)}" cy="{emu(bb.height)}"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
            f'<wp:wrapNone/><wp:docPr id="{inline.docPr.id}" name="{inline.docPr.name}"/><wp:cNvGraphicFramePr/></wp:anchor>')
        anchor.append(inline.graphic)
        inline.getparent().replace(inline, anchor)


def build_document(doc, idxs, ocr_pages, step=None):
    """把 doc 的 idxs 页转成 python-docx Document。返回 (Document, 提示列表)。"""
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.shared import Pt

    d = Document()
    d.styles["Normal"].paragraph_format.space_after = Pt(0)
    _modern_layout(d)
    fonts = Fonts(doc)
    warnings: list = []
    for k, i in enumerate(idxs):
        if step:
            step(k / len(idxs), f"转换第 {i + 1} 页…")
        page = doc[i]
        W, H = page.rect.width, page.rect.height
        scale = min(1.0, 1584 / max(W, H))  # Word 页面最大 22 英寸
        sec = d.sections[0] if k == 0 else d.add_section(WD_SECTION.NEW_PAGE)
        sec.page_width, sec.page_height = Pt(W * scale), Pt(H * scale)

        cache = {}

        def textless(page=page, i=i):
            if "img" not in cache:
                from converter import _render_without_text
                cache["img"] = _render_without_text(page.parent, i, 150)
            return cache["img"]

        flow, floating, notes = analyze_page(page, fonts, ocr_pages.get(i), textless)
        warnings += notes
        # 每个元素在 Word 里"期望的顶端位置"和高度。后面按这个绝对位置算段前距,偏差不会一路累积
        plan = []
        for e in flow:
            if e["kind"] == "para":
                _exact, top, h = para_metrics(e)
            elif e["kind"] == "table":
                top, h = e["bbox"].y0, e["bbox"].height
            else:
                top, h = e["bbox"].y0, e["bbox"].height + 2
            plan.append((e, top, h))
        if flow:
            l = min(e["bbox"].x0 for e in flow)
            r = W - max(e["bbox"].x1 for e in flow)
            t = min(top for _e, top, _h in plan)
        else:
            l = r = t = 54.0
        l, r = (min(max(v, 18.0), W / 3) for v in (l, r))
        r = max(r - 6.0, 12.0)  # 右边多留 6pt 余量:PDF 里排满一行的文字,Word 里不会因为差一点点就把最后一个字挤到下一行
        t = min(max(t, 18.0), H / 2)
        b = 18.0  # 底边留足余量:Word 排出来比 PDF 稍高一点时,不至于溢出到下一页
        sec.left_margin, sec.right_margin = Pt(l * scale), Pt(r * scale)
        sec.top_margin, sec.bottom_margin = Pt(t * scale), Pt(b * scale)
        sec.header_distance = sec.footer_distance = Pt(10)
        text_w = W - l - r
        write_floating(d, floating, scale, warnings)
        cursor = t + (1 if floating else 0)
        for e, top, h in plan:
            gap = max(0.0, top - cursor)
            if e["kind"] == "para":
                write_paragraph(d, page, fonts, e, l, text_w, W, gap, scale)
            elif e["kind"] == "table":
                if gap >= 2:
                    _spacer(d, gap * scale)
                write_table(d, page, fonts, e, text_w, scale)
            else:
                write_inline_image(d, e, l, text_w, W, gap, scale, warnings)
            cursor = max(cursor, top) + h
    return d, sorted(set(warnings))


def _modern_layout(d):
    """把文档设成 Word 2013+ 的排版模式。python-docx 的模板是旧兼容模式,两端对齐时不会压缩空格,
    而现在的 Word 导出的 PDF 里,很多行正是靠压缩空格才排下的,不改的话最后一个词会被挤到下一行。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    settings = d.settings.element
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        settings.append(compat)
    for cs in compat.findall(qn("w:compatSetting")):
        if cs.get(qn("w:name")) == "compatibilityMode":
            compat.remove(cs)
    cs = OxmlElement("w:compatSetting")
    cs.set(qn("w:name"), "compatibilityMode")
    cs.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
    cs.set(qn("w:val"), "15")
    compat.append(cs)


def _spacer(d, height_pt: float):
    """一个固定高度的空段落,用来在表格前撑出空白。"""
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Pt
    p = d.add_paragraph()
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p.paragraph_format.line_spacing = Pt(max(height_pt, 1))
    for r in p.runs:
        r.font.size = Pt(1)
